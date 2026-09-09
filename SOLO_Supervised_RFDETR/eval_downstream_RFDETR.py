#!/usr/bin/env python
"""Final locked downstream evaluation for the RF-DETR quality-assessment study.

The clinical workbook is treated as the final test set. Detector confidence
thresholds and all inference settings are loaded from ``qa_inference_RFDETR``
and verified against the independent calibration result before inference.
This script never selects or optimizes settings on the downstream labels.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import re
import sys
import time
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Sequence

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None  # type: ignore[assignment]

try:
    from PIL import ImageDraw, ImageFont
except Exception:
    ImageDraw = None  # type: ignore[assignment]
    ImageFont = None  # type: ignore[assignment]


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
for import_path in (SCRIPT_DIR, PROJECT_ROOT):
    import_string = str(import_path)
    if import_string not in sys.path:
        sys.path.insert(0, import_string)

import qa_inference_RFDETR as qa  # noqa: E402


DEFAULT_MANUAL_LABELS = Path(
    r"C:\Users\SH37YE\OneDrive\Full_FOV_Master_Review_2026-06-02_1154.xlsx"
)
DEFAULT_SHEET_NAME = "Master"
DEFAULT_OUTPUT_ROOT = PROJECT_ROOT / "DownstreamOutput_Downstream"
DEFAULT_CALIBRATION_SUMMARY = (
    PROJECT_ROOT
    / "EvaluationOutput"
    / "RFDETR_ThresholdCalibration_20260701-140419"
    / "threshold_calibration_summary.json"
)

EXPECTED_WORKBOOK_ROWS = 102
EXPECTED_INCLUDED_IMAGES = 98
EXPECTED_LABEL_COUNTS = {
    "Qualified": 42,
    "Partially Qualified": 29,
    "Not Qualified": 27,
}
EXPECTED_EXCLUSION_COUNTS = {
    "not_in_cvat_project": 1,
    "too_difficult_for_clinical_review": 3,
}
EXPECTED_MODEL_CLASS = "RFDETR2XLarge"
EXPECTED_MODEL_RESOLUTION = 880
EXPECTED_THRESHOLDS = {
    "Leucocyte": 0.36,
    "Squamous Epithelial Cell": 0.35,
}
EXPECTED_SLICE_SIZE = 640

LABELS_BY_ID = dict(qa.DOWNSTREAM_LABELS)
LABEL_IDS = sorted(LABELS_BY_ID)
LABEL_ID_BY_NORMALIZED_NAME = {
    re.sub(r"[^a-z0-9]+", "", name.casefold()): label_id
    for label_id, name in LABELS_BY_ID.items()
}


@dataclass(frozen=True)
class ClinicalRecord:
    source_row: int
    source_image_name: str
    manual_label_id: int
    manual_label: str
    epithelial_count_bin: str
    leucocyte_count_bin: str
    geckler_class: str
    collapsed_geckler_label: str
    murray_washington_label: str
    annotator: str
    comment: str
    image_path: str = ""


@dataclass(frozen=True)
class ExcludedRecord:
    source_row: int
    source_image_name: str
    source_label: str
    annotator: str
    comment: str
    exclusion_reason: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Evaluate the locked RF-DETR downstream quality classifier against "
            "the final expert-reviewed full-FOV workbook."
        )
    )
    parser.add_argument("--checkpoint", type=Path, default=qa.DEFAULT_CHECKPOINT)
    parser.add_argument("--images-root", type=Path, default=Path(qa.DEFAULT_IMAGES_ROOT))
    parser.add_argument("--manual-labels", type=Path, default=DEFAULT_MANUAL_LABELS)
    parser.add_argument("--sheet-name", default=DEFAULT_SHEET_NAME)
    parser.add_argument("--calibration-summary", type=Path, default=DEFAULT_CALIBRATION_SUMMARY)
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument(
        "--preflight-only",
        action="store_true",
        help="Validate every input and locked setting without loading the detector or writing outputs.",
    )
    parser.add_argument(
        "--no-overlays",
        action="store_true",
        help="Do not save annotated full-FOV images. Metrics and tabular results are still saved.",
    )
    return parser.parse_args()


def normalized_label(value: Any) -> int:
    text = str(value or "").strip()
    normalized = re.sub(r"[^a-z0-9]+", "", text.casefold())
    if normalized in LABEL_ID_BY_NORMALIZED_NAME:
        return LABEL_ID_BY_NORMALIZED_NAME[normalized]
    raise ValueError(f"Unsupported clinical label: {value!r}")


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


COUNT_BINS = ("0-9", "10-25", "26+")


def normalized_count_bin(value: Any) -> str:
    """Normalize the expert count categories without inventing exact counts."""
    text = _cell_text(value).casefold().replace("–", "-").replace("—", "-")
    compact = re.sub(r"\s+", "", text.replace("til", "-"))
    if compact in {"0-9", "0to9"}:
        return "0-9"
    if compact in {"10-25", "10to25"}:
        return "10-25"
    if compact in {"26+", ">25", "26"}:
        return "26+"
    try:
        number = int(float(compact))
    except (TypeError, ValueError):
        raise ValueError(f"Unsupported expert count category: {value!r}")
    if number <= 9:
        return "0-9"
    if number <= 25:
        return "10-25"
    return "26+"


def count_bin_from_integer(value: int) -> str:
    if value <= 9:
        return "0-9"
    if value <= 25:
        return "10-25"
    return "26+"


def geckler_class(epithelial_bin: str, leucocyte_bin: str) -> str:
    """Return Geckler group 1-6 from the study's three expert count bins."""
    if epithelial_bin == "26+":
        return {"0-9": "G1", "10-25": "G2", "26+": "G3"}[leucocyte_bin]
    if epithelial_bin == "10-25" and leucocyte_bin == "26+":
        return "G4"
    if epithelial_bin == "0-9" and leucocyte_bin == "26+":
        return "G5"
    return "G6"


def collapsed_geckler_label(group: str) -> str:
    """Collapse Geckler to acceptable, unacceptable, and unknown."""
    if group in {"G4", "G5"}:
        return "Acceptable"
    if group in {"G1", "G2", "G3"}:
        return "Unacceptable"
    return "Unknown"


def murray_washington_label(epithelial_bin: str, leucocyte_bin: str) -> str:
    """Binary Murray-Washington culture-quality interpretation."""
    return (
        "Acceptable"
        if epithelial_bin == "0-9" and leucocyte_bin == "26+"
        else "Unacceptable"
    )


def read_clinical_workbook(
    workbook_path: Path,
    sheet_name: str,
) -> tuple[list[ClinicalRecord], list[ExcludedRecord], int]:
    if load_workbook is None:
        raise ImportError("openpyxl is required to read the final clinical workbook.")
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Clinical workbook does not exist: {workbook_path}")

    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except PermissionError as exc:
        raise PermissionError(
            f"Cannot read the clinical workbook: {workbook_path}. "
            "Save and close it in Excel, then run the command again."
        ) from exc

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet {sheet_name!r} was not found. Available sheets: "
                f"{', '.join(workbook.sheetnames)}"
            )
        sheet = workbook[sheet_name]
        header = [_cell_text(cell.value).casefold() for cell in sheet[1][:6]]
        expected_header = [
            "image", "tag", "epithelial cells", "leucocytes", "annotator", "comments"
        ]
        if [value.rstrip() for value in header[:6]] != expected_header:
            raise ValueError(
                f"Expected columns A:F in {sheet_name!r} to be {expected_header!r}; "
                f"found {header[:6]!r}."
            )

        included: list[ClinicalRecord] = []
        excluded: list[ExcludedRecord] = []
        workbook_rows = 0

        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, min_col=1, max_col=6, values_only=True),
            start=2,
        ):
            (
                image_name,
                label_value,
                epithelial_value,
                leucocyte_value,
                annotator_value,
                comment_value,
            ) = row
            if all(value is None or not str(value).strip() for value in row):
                continue

            workbook_rows += 1
            image_text = _cell_text(image_name)
            label_text = _cell_text(label_value)
            annotator = _cell_text(annotator_value)
            comment = _cell_text(comment_value)

            if not image_text:
                raise ValueError(f"Workbook row {row_number} has no image name.")

            if label_text.casefold() == "er ikke i projektet":
                excluded.append(
                    ExcludedRecord(
                        source_row=row_number,
                        source_image_name=image_text,
                        source_label=label_text,
                        annotator=annotator,
                        comment=comment,
                        exclusion_reason="not_in_cvat_project",
                    )
                )
                continue

            if not label_text:
                excluded.append(
                    ExcludedRecord(
                        source_row=row_number,
                        source_image_name=image_text,
                        source_label="",
                        annotator=annotator,
                        comment=comment,
                        exclusion_reason="too_difficult_for_clinical_review",
                    )
                )
                continue

            try:
                label_id = normalized_label(label_text)
                epithelial_bin = normalized_count_bin(epithelial_value)
                leucocyte_bin = normalized_count_bin(leucocyte_value)
            except ValueError as exc:
                raise ValueError(f"Workbook row {row_number}: {exc}") from exc

            included.append(
                ClinicalRecord(
                    source_row=row_number,
                    source_image_name=image_text,
                    manual_label_id=label_id,
                    manual_label=LABELS_BY_ID[label_id],
                    epithelial_count_bin=epithelial_bin,
                    leucocyte_count_bin=leucocyte_bin,
                    geckler_class=geckler_class(epithelial_bin, leucocyte_bin),
                    collapsed_geckler_label=collapsed_geckler_label(
                        geckler_class(epithelial_bin, leucocyte_bin)
                    ),
                    murray_washington_label=murray_washington_label(
                        epithelial_bin, leucocyte_bin
                    ),
                    annotator=annotator,
                    comment=comment,
                )
            )
    finally:
        workbook.close()

    return included, excluded, workbook_rows


def build_full_fov_index(images_root: Path) -> dict[str, list[Path]]:
    """Index only full-FOV files directly inside each Sample directory.

    Patch subdirectories are deliberately not traversed, so a clinical full-FOV
    name can never resolve to one of the detector-training patches.
    """
    if not images_root.is_dir():
        raise FileNotFoundError(f"Full-FOV image root does not exist: {images_root}")

    index: dict[str, list[Path]] = defaultdict(list)
    for sample_dir in images_root.iterdir():
        if not sample_dir.is_dir():
            continue
        for candidate in sample_dir.iterdir():
            if candidate.is_file() and candidate.suffix.casefold() in qa.IMAGE_EXTENSIONS:
                index[candidate.name.casefold()].append(candidate)
    return dict(index)


def resolve_clinical_images(
    records: Sequence[ClinicalRecord],
    images_root: Path,
) -> list[ClinicalRecord]:
    index = build_full_fov_index(images_root)
    resolved: list[ClinicalRecord] = []
    missing: list[str] = []
    ambiguous: dict[str, list[str]] = {}

    for record in records:
        candidates = index.get(Path(record.source_image_name).name.casefold(), [])
        if not candidates:
            missing.append(record.source_image_name)
            continue
        if len(candidates) > 1:
            ambiguous[record.source_image_name] = [str(path) for path in candidates]
            continue
        resolved.append(
            ClinicalRecord(
                **{
                    **asdict(record),
                    "image_path": str(candidates[0].resolve()),
                }
            )
        )

    if missing or ambiguous:
        messages: list[str] = []
        if missing:
            messages.append("missing=" + ", ".join(sorted(missing)))
        if ambiguous:
            messages.append(
                "ambiguous="
                + "; ".join(
                    f"{name}: {paths}" for name, paths in sorted(ambiguous.items())
                )
            )
        raise ValueError("Clinical image resolution failed: " + " | ".join(messages))
    return resolved


def _same_path(first: Path, second: Path) -> bool:
    return os.path.normcase(os.path.abspath(first)) == os.path.normcase(os.path.abspath(second))


def _assert_close(actual: float, expected: float, description: str) -> None:
    if not math.isclose(float(actual), float(expected), rel_tol=0.0, abs_tol=1e-6):
        raise ValueError(f"{description} is {actual!r}; expected {expected!r}.")


def verify_calibration(
    calibration_summary_path: Path,
    checkpoint: Path,
) -> dict[str, Any]:
    if not calibration_summary_path.is_file():
        raise FileNotFoundError(
            f"Calibration summary does not exist: {calibration_summary_path}"
        )
    with calibration_summary_path.open("r", encoding="utf-8") as handle:
        summary = json.load(handle)

    calibration_checkpoint = Path(str(summary.get("checkpoint", "")))
    if not _same_path(calibration_checkpoint, checkpoint):
        raise ValueError(
            "The calibration checkpoint differs from the downstream checkpoint: "
            f"{calibration_checkpoint} != {checkpoint}"
        )
    if summary.get("model_class") != EXPECTED_MODEL_CLASS:
        raise ValueError(
            f"Calibration model class is {summary.get('model_class')!r}; "
            f"expected {EXPECTED_MODEL_CLASS!r}."
        )

    selected = summary.get("selected_thresholds", {})
    for class_name, expected in EXPECTED_THRESHOLDS.items():
        if class_name not in selected:
            raise ValueError(f"Calibration summary has no threshold for {class_name!r}.")
        _assert_close(selected[class_name], expected, f"Calibrated {class_name} threshold")

    return {
        "summary_path": str(calibration_summary_path.resolve()),
        "timestamp": summary.get("timestamp"),
        "calibration_root": summary.get("calibration_root"),
        "image_count": summary.get("image_count"),
        "selection_metric": summary.get("selection_metric"),
        "selected_thresholds": selected,
        "checkpoint": str(calibration_checkpoint),
    }


def verify_locked_inference(
    checkpoint: Path,
) -> tuple[str, int | None, list[str], dict[str, float], dict[str, Any]]:
    model_class, model_resolution, class_names, thresholds = qa.inspect_checkpoint(checkpoint)
    settings = qa.downstream_inference_settings()

    if model_class != EXPECTED_MODEL_CLASS:
        raise ValueError(
            f"Downstream model class is {model_class!r}; expected {EXPECTED_MODEL_CLASS!r}."
        )
    if model_resolution != EXPECTED_MODEL_RESOLUTION:
        raise ValueError(
            f"Downstream model resolution is {model_resolution!r}; "
            f"expected {EXPECTED_MODEL_RESOLUTION}."
        )
    if class_names != list(EXPECTED_THRESHOLDS):
        raise ValueError(
            f"Downstream classes are {class_names!r}; expected {list(EXPECTED_THRESHOLDS)!r}."
        )
    for class_name, expected in EXPECTED_THRESHOLDS.items():
        _assert_close(thresholds[class_name], expected, f"Runtime {class_name} threshold")
        _assert_close(
            settings["class_score_thresholds"][class_name],
            expected,
            f"Shared {class_name} threshold",
        )
    if (
        settings["slice_height"] != EXPECTED_SLICE_SIZE
        or settings["slice_width"] != EXPECTED_SLICE_SIZE
    ):
        raise ValueError(
            "Downstream SAHI slices must be "
            f"{EXPECTED_SLICE_SIZE}x{EXPECTED_SLICE_SIZE}; found "
            f"{settings['slice_width']}x{settings['slice_height']}."
        )

    return model_class, model_resolution, class_names, thresholds, settings


def verify_clinical_dataset(
    included: Sequence[ClinicalRecord],
    excluded: Sequence[ExcludedRecord],
    workbook_rows: int,
) -> None:
    label_counts = Counter(record.manual_label for record in included)
    exclusion_counts = Counter(record.exclusion_reason for record in excluded)

    if workbook_rows != EXPECTED_WORKBOOK_ROWS:
        raise ValueError(
            f"Clinical workbook contains {workbook_rows} non-header rows; "
            f"expected {EXPECTED_WORKBOOK_ROWS}."
        )
    if len(included) != EXPECTED_INCLUDED_IMAGES:
        raise ValueError(
            f"Clinical test set contains {len(included)} included images; "
            f"expected {EXPECTED_INCLUDED_IMAGES}."
        )
    if dict(label_counts) != EXPECTED_LABEL_COUNTS:
        raise ValueError(
            f"Clinical label counts are {dict(label_counts)!r}; "
            f"expected {EXPECTED_LABEL_COUNTS!r}."
        )
    if dict(exclusion_counts) != EXPECTED_EXCLUSION_COUNTS:
        raise ValueError(
            f"Clinical exclusion counts are {dict(exclusion_counts)!r}; "
            f"expected {EXPECTED_EXCLUSION_COUNTS!r}."
        )


def run_preflight(args: argparse.Namespace) -> dict[str, Any]:
    checkpoint = args.checkpoint.expanduser()
    images_root = args.images_root.expanduser()
    workbook_path = args.manual_labels.expanduser()
    calibration_path = args.calibration_summary.expanduser()

    if not checkpoint.is_file():
        raise FileNotFoundError(f"Checkpoint does not exist: {checkpoint}")

    included, excluded, workbook_rows = read_clinical_workbook(
        workbook_path,
        args.sheet_name,
    )
    verify_clinical_dataset(included, excluded, workbook_rows)
    resolved = resolve_clinical_images(included, images_root)
    model_class, model_resolution, class_names, thresholds, settings = (
        verify_locked_inference(checkpoint)
    )
    calibration = verify_calibration(calibration_path, checkpoint)

    return {
        "checkpoint": str(checkpoint.resolve()),
        "model_class": model_class,
        "model_resolution": model_resolution,
        "class_names": class_names,
        "inference_settings": settings,
        "runtime_thresholds": thresholds,
        "clinical_workbook": str(workbook_path.resolve()),
        "clinical_sheet": args.sheet_name,
        "workbook_rows": workbook_rows,
        "included_image_count": len(resolved),
        "manual_label_counts": dict(
            Counter(record.manual_label for record in resolved)
        ),
        "expert_geckler_counts": dict(
            Counter(record.geckler_class for record in resolved)
        ),
        "expert_murray_washington_counts": dict(
            Counter(record.murray_washington_label for record in resolved)
        ),
        "expert_collapsed_geckler_counts": dict(
            Counter(record.collapsed_geckler_label for record in resolved)
        ),
        "excluded_image_count": len(excluded),
        "exclusion_counts": dict(
            Counter(record.exclusion_reason for record in excluded)
        ),
        "images_root": str(images_root.resolve()),
        "calibration": calibration,
        "_resolved_records": resolved,
        "_excluded_records": list(excluded),
    }


def print_preflight(summary: dict[str, Any]) -> None:
    settings = summary["inference_settings"]
    print("Preflight passed.")
    print(
        f"Model: {summary['model_class']} at internal resolution "
        f"{summary['model_resolution']} ({summary['checkpoint']})"
    )
    print(
        "SAHI: "
        f"{settings['slice_width']}x{settings['slice_height']} slices, "
        f"{settings['overlap_width_ratio']:.0%} overlap, "
        f"{settings['postprocess']['type']} "
        f"{settings['postprocess']['match_metric']}="
        f"{settings['postprocess']['match_threshold']:.2f}"
    )
    thresholds = settings["class_score_thresholds"]
    print(
        "Calibrated thresholds: "
        + ", ".join(f"{name}={value:.2f}" for name, value in thresholds.items())
    )
    print(
        f"Clinical set: {summary['included_image_count']} included / "
        f"{summary['workbook_rows']} reviewed rows; "
        f"{summary['excluded_image_count']} excluded"
    )
    print(f"Manual labels: {summary['manual_label_counts']}")
    print(f"Resolved full-FOV root: {summary['images_root']}")


def create_output_dir(output_root: Path) -> Path:
    output_root.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    output_dir = output_root / timestamp
    suffix = 2
    while output_dir.exists():
        output_dir = output_root / f"{timestamp}_{suffix:02d}"
        suffix += 1
    output_dir.mkdir(parents=False)
    return output_dir


def write_csv(
    path: Path,
    rows: Iterable[dict[str, Any]],
    fieldnames: Sequence[str],
) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def confusion_matrix_counts(
    true_ids: Sequence[int],
    predicted_ids: Sequence[int],
) -> list[list[int]]:
    matrix = [[0 for _ in LABEL_IDS] for _ in LABEL_IDS]
    positions = {label_id: index for index, label_id in enumerate(LABEL_IDS)}
    for true_id, predicted_id in zip(true_ids, predicted_ids):
        matrix[positions[true_id]][positions[predicted_id]] += 1
    return matrix


def safe_ratio(numerator: float, denominator: float) -> float:
    return float(numerator / denominator) if denominator else 0.0


def calculate_metrics(
    true_ids: Sequence[int],
    predicted_ids: Sequence[int],
) -> tuple[dict[str, Any], list[dict[str, Any]], list[list[int]]]:
    if len(true_ids) != len(predicted_ids) or not true_ids:
        raise ValueError("Metric inputs must be non-empty and have equal lengths.")

    matrix = confusion_matrix_counts(true_ids, predicted_ids)
    total = len(true_ids)
    per_class: list[dict[str, Any]] = []
    for index, label_id in enumerate(LABEL_IDS):
        tp = matrix[index][index]
        fn = sum(matrix[index]) - tp
        fp = sum(row[index] for row in matrix) - tp
        tn = total - tp - fn - fp
        precision = safe_ratio(tp, tp + fp)
        recall = safe_ratio(tp, tp + fn)
        f1 = safe_ratio(2 * precision * recall, precision + recall)
        per_class.append(
            {
                "label_id": label_id,
                "label": LABELS_BY_ID[label_id],
                "support": tp + fn,
                "tp": tp,
                "fp": fp,
                "fn": fn,
                "tn": tn,
                "precision": precision,
                "recall": recall,
                "f1": f1,
                "specificity": safe_ratio(tn, tn + fp),
            }
        )

    accuracy = safe_ratio(sum(matrix[i][i] for i in range(len(LABEL_IDS))), total)
    macro_precision = sum(row["precision"] for row in per_class) / len(per_class)
    macro_recall = sum(row["recall"] for row in per_class) / len(per_class)
    macro_f1 = sum(row["f1"] for row in per_class) / len(per_class)
    weighted_precision = sum(row["precision"] * row["support"] for row in per_class) / total
    weighted_recall = sum(row["recall"] * row["support"] for row in per_class) / total
    weighted_f1 = sum(row["f1"] * row["support"] for row in per_class) / total
    mean_absolute_class_error = sum(
        abs(true_id - predicted_id)
        for true_id, predicted_id in zip(true_ids, predicted_ids)
    ) / total
    within_one_class_accuracy = sum(
        abs(true_id - predicted_id) <= 1
        for true_id, predicted_id in zip(true_ids, predicted_ids)
    ) / total

    true_counts = Counter(true_ids)
    predicted_counts = Counter(predicted_ids)
    expected_agreement = sum(
        true_counts[label_id] * predicted_counts[label_id] for label_id in LABEL_IDS
    ) / (total * total)
    cohen_kappa = safe_ratio(accuracy - expected_agreement, 1.0 - expected_agreement)

    metrics = {
        "n_images": total,
        "accuracy": accuracy,
        "balanced_accuracy": macro_recall,
        "macro_precision": macro_precision,
        "macro_recall": macro_recall,
        "macro_f1": macro_f1,
        "weighted_precision": weighted_precision,
        "weighted_recall": weighted_recall,
        "weighted_f1": weighted_f1,
        "cohen_kappa": cohen_kappa,
        "mean_absolute_class_error": mean_absolute_class_error,
        "within_one_class_accuracy": within_one_class_accuracy,
    }
    return metrics, per_class, matrix


def calculate_named_metrics(
    true_labels: Sequence[str],
    predicted_labels: Sequence[str],
    labels: Sequence[str],
) -> tuple[dict[str, Any], list[dict[str, Any]], list[list[int]]]:
    """Calculate the same classification metrics for an arbitrary label scheme."""
    if len(true_labels) != len(predicted_labels) or not true_labels:
        raise ValueError("Metric inputs must be non-empty and have equal lengths.")
    positions = {label: index for index, label in enumerate(labels)}
    unknown = (set(true_labels) | set(predicted_labels)) - set(labels)
    if unknown:
        raise ValueError(f"Unexpected classification labels: {sorted(unknown)!r}")
    matrix = [[0 for _ in labels] for _ in labels]
    for true_label, predicted_label in zip(true_labels, predicted_labels):
        matrix[positions[true_label]][positions[predicted_label]] += 1

    total = len(true_labels)
    per_class: list[dict[str, Any]] = []
    for index, label in enumerate(labels):
        tp = matrix[index][index]
        fn = sum(matrix[index]) - tp
        fp = sum(row[index] for row in matrix) - tp
        tn = total - tp - fn - fp
        precision = safe_ratio(tp, tp + fp)
        recall = safe_ratio(tp, tp + fn)
        per_class.append(
            {
                "label": label,
                "support": tp + fn,
                "tp": tp,
                "fp": fp,
                "fn": fn,
                "tn": tn,
                "precision": precision,
                "recall": recall,
                "f1": safe_ratio(2 * precision * recall, precision + recall),
                "specificity": safe_ratio(tn, tn + fp),
            }
        )

    accuracy = safe_ratio(sum(matrix[i][i] for i in range(len(labels))), total)
    macro_recall = sum(row["recall"] for row in per_class) / len(per_class)
    macro_f1 = sum(row["f1"] for row in per_class) / len(per_class)
    true_counts = Counter(true_labels)
    predicted_counts = Counter(predicted_labels)
    expected = sum(
        true_counts[label] * predicted_counts[label] for label in labels
    ) / (total * total)
    return (
        {
            "n_images": total,
            "accuracy": accuracy,
            "balanced_accuracy": macro_recall,
            "macro_f1": macro_f1,
            "cohen_kappa": safe_ratio(accuracy - expected, 1.0 - expected),
        },
        per_class,
        matrix,
    )


def save_named_confusion_csv(
    path: Path,
    matrix: Sequence[Sequence[int]],
    labels: Sequence[str],
) -> None:
    fieldnames = ["reference_label"] + [f"predicted_{label}" for label in labels]
    rows = []
    for label, values in zip(labels, matrix):
        row: dict[str, Any] = {"reference_label": label}
        row.update(
            {f"predicted_{predicted}": value for predicted, value in zip(labels, values)}
        )
        rows.append(row)
    write_csv(path, rows, fieldnames)


def save_confusion_csv(path: Path, matrix: Sequence[Sequence[int]]) -> None:
    fieldnames = ["manual_label"] + [
        f"predicted_{LABELS_BY_ID[label_id]}" for label_id in LABEL_IDS
    ]
    rows = []
    for label_id, values in zip(LABEL_IDS, matrix):
        row: dict[str, Any] = {"manual_label": LABELS_BY_ID[label_id]}
        row.update(
            {
                f"predicted_{LABELS_BY_ID[predicted_id]}": value
                for predicted_id, value in zip(LABEL_IDS, values)
            }
        )
        rows.append(row)
    write_csv(path, rows, fieldnames)


def save_confusion_figure(path: Path, matrix: Sequence[Sequence[int]]) -> bool:
    try:
        import matplotlib.pyplot as plt
    except Exception:
        return False

    figure, axis = plt.subplots(figsize=(8, 7))
    image = axis.imshow(matrix, cmap="Blues")
    figure.colorbar(image, ax=axis, fraction=0.046, pad=0.04)
    names = [LABELS_BY_ID[label_id] for label_id in LABEL_IDS]
    axis.set(
        xticks=range(len(names)),
        yticks=range(len(names)),
        xticklabels=names,
        yticklabels=names,
        xlabel="Predicted label",
        ylabel="Clinical label",
        title="Final downstream confusion matrix",
    )
    plt.setp(axis.get_xticklabels(), rotation=25, ha="right")
    maximum = max(max(row) for row in matrix) if matrix else 0
    for row_index, row in enumerate(matrix):
        for column_index, value in enumerate(row):
            axis.text(
                column_index,
                row_index,
                str(value),
                ha="center",
                va="center",
                color="white" if maximum and value > maximum / 2 else "black",
            )
    figure.tight_layout()
    figure.savefig(path, dpi=200, bbox_inches="tight")
    plt.close(figure)
    return True


def add_clinical_label_to_overlay(
    image: Any,
    manual_label: str,
    predicted_label: str,
) -> None:
    if ImageDraw is None or ImageFont is None:
        return
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default()
    text = f"Clinical: {manual_label} | Predicted: {predicted_label}"
    try:
        bounds = draw.textbbox((0, 0), text, font=font)
        text_width = bounds[2] - bounds[0]
        text_height = bounds[3] - bounds[1]
    except Exception:
        text_width = len(text) * 7
        text_height = 12
    padding = 6
    left = max(0, image.width - text_width - 2 * padding - 8)
    top = 8
    draw.rectangle(
        [left, top, image.width - 8, top + text_height + 2 * padding],
        fill=(0, 0, 0),
    )
    draw.text(
        (left + padding, top + padding),
        text,
        fill=(255, 255, 255),
        font=font,
    )


def evaluate(args: argparse.Namespace, preflight: dict[str, Any]) -> Path:
    output_dir = create_output_dir(args.output_root.expanduser())
    overlays_dir = output_dir / "overlays"
    if not args.no_overlays:
        overlays_dir.mkdir()

    resolved: list[ClinicalRecord] = preflight.pop("_resolved_records")
    excluded: list[ExcludedRecord] = preflight.pop("_excluded_records")

    print(f"Loading detector once for {len(resolved)} final full-FOV images...")
    runtime = qa.build_runtime_for_checkpoint(args.checkpoint.expanduser())
    rows: list[dict[str, Any]] = []
    evaluation_started = time.perf_counter()

    for image_index, record in enumerate(resolved, start=1):
        image_path = Path(record.image_path)
        started = time.perf_counter()
        result = qa.run_inference_on_image(image_path, runtime)
        inference_seconds = time.perf_counter() - started

        row = {
            **asdict(record),
            "predicted_label_id": result.predicted_label_id,
            "predicted_label": result.predicted_label,
            "predicted_epithelial_count_bin": count_bin_from_integer(
                result.n_squamous_epithelial_cell
            ),
            "predicted_leucocyte_count_bin": count_bin_from_integer(
                result.n_leucocyte
            ),
            "correct": int(result.predicted_label_id == record.manual_label_id),
            "absolute_class_error": abs(
                result.predicted_label_id - record.manual_label_id
            ),
            "n_leucocyte": result.n_leucocyte,
            "n_squamous_epithelial_cell": result.n_squamous_epithelial_cell,
            "leucocyte_score": result.leucocyte_score,
            "squamous_epithelial_score": result.squamous_epithelial_score,
            "total_quality_score": result.total_quality_score,
            "n_predictions_raw": result.n_predictions_raw,
            "n_predictions_kept_before_duplicate_suppression": (
                result.n_predictions_kept_before_duplicate_suppression
            ),
            "n_cross_class_duplicates_suppressed": (
                result.n_cross_class_duplicates_suppressed
            ),
            "n_predictions_kept": result.n_predictions_kept,
            "inference_seconds": inference_seconds,
        }
        row["predicted_geckler_class"] = geckler_class(
            row["predicted_epithelial_count_bin"],
            row["predicted_leucocyte_count_bin"],
        )
        row["predicted_collapsed_geckler_label"] = collapsed_geckler_label(
            row["predicted_geckler_class"]
        )
        row["predicted_murray_washington_label"] = murray_washington_label(
            row["predicted_epithelial_count_bin"],
            row["predicted_leucocyte_count_bin"],
        )
        rows.append(row)

        if not args.no_overlays:
            overlay = qa.render_result_overlay(
                image_path,
                result,
                runtime.class_names,
            )
            add_clinical_label_to_overlay(
                overlay,
                record.manual_label,
                result.predicted_label,
            )
            overlay_name = (
                f"{image_index:03d}_row{record.source_row}_"
                f"{image_path.stem}.jpg"
            )
            overlay.save(overlays_dir / overlay_name, quality=92)
            overlay.close()

        print(
            f"[{image_index:03d}/{len(resolved):03d}] {image_path.name}: "
            f"clinical={record.manual_label}, predicted={result.predicted_label}, "
            f"Leu={result.n_leucocyte}, Epi={result.n_squamous_epithelial_cell}, "
            f"{inference_seconds:.1f}s"
        )

    elapsed_seconds = time.perf_counter() - evaluation_started
    true_ids = [int(row["manual_label_id"]) for row in rows]
    predicted_ids = [int(row["predicted_label_id"]) for row in rows]
    metrics, per_class_metrics, matrix = calculate_metrics(true_ids, predicted_ids)
    scheme_inputs = {
        "geckler": (
            [str(row["geckler_class"]) for row in rows],
            [str(row["predicted_geckler_class"]) for row in rows],
            [f"G{index}" for index in range(1, 7)],
        ),
        "murray_washington": (
            [str(row["murray_washington_label"]) for row in rows],
            [str(row["predicted_murray_washington_label"]) for row in rows],
            ["Acceptable", "Unacceptable"],
        ),
        "collapsed_geckler": (
            [str(row["collapsed_geckler_label"]) for row in rows],
            [str(row["predicted_collapsed_geckler_label"]) for row in rows],
            ["Acceptable", "Unacceptable", "Unknown"],
        ),
    }
    scheme_results: dict[str, Any] = {}
    for scheme_name, (true_labels, predicted_labels, scheme_labels) in scheme_inputs.items():
        scheme_metrics, scheme_per_class, scheme_matrix = calculate_named_metrics(
            true_labels, predicted_labels, scheme_labels
        )
        scheme_csv = output_dir / f"downstream_{scheme_name}_confusion_matrix.csv"
        save_named_confusion_csv(scheme_csv, scheme_matrix, scheme_labels)
        scheme_results[scheme_name] = {
            "labels": scheme_labels,
            "reference_label_counts": dict(Counter(true_labels)),
            "predicted_label_counts": dict(Counter(predicted_labels)),
            "metrics": scheme_metrics,
            "per_class_metrics": scheme_per_class,
            "confusion_matrix": scheme_matrix,
            "confusion_matrix_csv": str(scheme_csv.resolve()),
        }

    prediction_fields = list(rows[0])
    predictions_path = output_dir / "downstream_predictions.csv"
    write_csv(predictions_path, rows, prediction_fields)

    exclusions_path = output_dir / "downstream_exclusions.csv"
    exclusion_rows = [asdict(record) for record in excluded]
    write_csv(
        exclusions_path,
        exclusion_rows,
        list(ExcludedRecord.__dataclass_fields__),
    )

    per_class_path = output_dir / "downstream_per_class_metrics.csv"
    write_csv(per_class_path, per_class_metrics, list(per_class_metrics[0]))

    confusion_csv_path = output_dir / "downstream_confusion_matrix.csv"
    save_confusion_csv(confusion_csv_path, matrix)
    confusion_png_path = output_dir / "downstream_confusion_matrix.png"
    confusion_figure_saved = save_confusion_figure(confusion_png_path, matrix)

    summary = {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "purpose": (
            "Final downstream evaluation using settings locked before inspecting "
            "downstream performance."
        ),
        "preflight": preflight,
        "dataset": {
            "included_image_count": len(rows),
            "excluded_image_count": len(excluded),
            "manual_label_counts": dict(
                Counter(row["manual_label"] for row in rows)
            ),
            "predicted_label_counts": dict(
                Counter(row["predicted_label"] for row in rows)
            ),
            "exclusions": exclusion_rows,
        },
        "metrics": metrics,
        "per_class_metrics": per_class_metrics,
        "classification_schemes": scheme_results,
        "confusion_matrix": {
            "row_axis": "clinical_label",
            "column_axis": "predicted_label",
            "labels": [LABELS_BY_ID[label_id] for label_id in LABEL_IDS],
            "counts": matrix,
        },
        "runtime": {
            "total_evaluation_seconds": elapsed_seconds,
            "mean_inference_seconds_per_image": safe_ratio(
                sum(float(row["inference_seconds"]) for row in rows),
                len(rows),
            ),
        },
        "outputs": {
            "predictions_csv": str(predictions_path.resolve()),
            "exclusions_csv": str(exclusions_path.resolve()),
            "per_class_metrics_csv": str(per_class_path.resolve()),
            "confusion_matrix_csv": str(confusion_csv_path.resolve()),
            "confusion_matrix_png": (
                str(confusion_png_path.resolve())
                if confusion_figure_saved
                else None
            ),
            "overlays_dir": (
                str(overlays_dir.resolve()) if not args.no_overlays else None
            ),
        },
    }
    summary_path = output_dir / "downstream_evaluation_summary.json"
    with summary_path.open("w", encoding="utf-8") as handle:
        json.dump(summary, handle, indent=2, ensure_ascii=False)

    print(f"Final downstream evaluation complete: {output_dir}")
    print(
        f"Accuracy={metrics['accuracy']:.4f}, "
        f"balanced accuracy={metrics['balanced_accuracy']:.4f}, "
        f"macro F1={metrics['macro_f1']:.4f}"
    )
    return output_dir


def main() -> int:
    args = parse_args()
    try:
        preflight = run_preflight(args)
        print_preflight(preflight)
        if args.preflight_only:
            return 0
        evaluate(args, preflight)
        return 0
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
