#!/usr/bin/env python
"""Build one Word report covering detector, calibration, and downstream results.

The report is generated from saved evaluation artifacts. No model inference or
threshold optimization is performed here. Microsoft Word is automated through
the locally installed COM interface so no additional document package is
required in the ai_powmic environment.
"""

from __future__ import annotations

import argparse
import contextlib
import csv
import io
import json
import re
import sys
import hashlib
import math
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Sequence

import numpy as np
from openpyxl import load_workbook


def accuracy_score(true: Sequence[Any], predicted: Sequence[Any]) -> float:
    return float(np.mean(np.asarray(true) == np.asarray(predicted)))


def balanced_accuracy_score(true: Sequence[Any], predicted: Sequence[Any]) -> float:
    true_array = np.asarray(true)
    predicted_array = np.asarray(predicted)
    labels = list(dict.fromkeys(true))
    recalls = [
        float(np.mean(predicted_array[true_array == label] == label))
        for label in labels
    ]
    return float(np.mean(recalls))


def f1_score(
    true: Sequence[Any], predicted: Sequence[Any], *, labels: Sequence[Any],
    average: str, zero_division: int = 0,
) -> float:
    if average != "macro":
        raise ValueError("Only macro F1 is implemented in this report builder.")
    scores = []
    for label in labels:
        tp = sum(a == label and b == label for a, b in zip(true, predicted))
        fp = sum(a != label and b == label for a, b in zip(true, predicted))
        fn = sum(a == label and b != label for a, b in zip(true, predicted))
        precision = tp / (tp + fp) if tp + fp else float(zero_division)
        recall = tp / (tp + fn) if tp + fn else float(zero_division)
        scores.append(
            2 * precision * recall / (precision + recall)
            if precision + recall else float(zero_division)
        )
    return float(np.mean(scores))


def cohen_kappa_score(
    true: Sequence[Any], predicted: Sequence[Any], labels: Sequence[Any] | None = None,
    weights: str | None = None,
) -> float:
    labels = list(labels or sorted(set([*true, *predicted])))
    total = len(true)
    if weights == "quadratic":
        positions = {label: index for index, label in enumerate(labels)}
        denominator_scale = max(1, len(labels) - 1) ** 2
        observed = sum(
            (positions[a] - positions[b]) ** 2 / denominator_scale
            for a, b in zip(true, predicted)
        ) / total
        true_counts = Counter(true)
        predicted_counts = Counter(predicted)
        expected = sum(
            true_counts[a] * predicted_counts[b]
            * ((positions[a] - positions[b]) ** 2 / denominator_scale)
            for a in labels for b in labels
        ) / (total * total)
        return float(1 - observed / expected) if expected else 0.0
    if weights is not None:
        raise ValueError(f"Unsupported kappa weights: {weights!r}")
    observed = accuracy_score(true, predicted)
    true_counts = Counter(true)
    predicted_counts = Counter(predicted)
    expected = sum(true_counts[label] * predicted_counts[label] for label in labels) / (total * total)
    return float((observed - expected) / (1 - expected)) if expected != 1 else 0.0

try:
    import win32com.client
except Exception:  # pragma: no cover - optional native Word backend
    win32com = None  # type: ignore[assignment]


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
OBJECT_OUTPUT = (
    PROJECT_ROOT
    / "EvaluationOutput"
    / "TwoClass_HPO009_EMA_TestEval_20260701-140821"
)
CALIBRATION_OUTPUT = (
    PROJECT_ROOT
    / "EvaluationOutput"
    / "RFDETR_ThresholdCalibration_20260701-140419"
)
DOWNSTREAM_ROOT = PROJECT_ROOT / "DownstreamOutput_Downstream"
SPLIT_SUMMARY = (
    SCRIPT_DIR
    / "Stat_Dataset"
    / "QA-2025v1_TwoClass_OVR_V2_20260618-101346"
    / "split_summary.json"
)
TEST_COCO = SPLIT_SUMMARY.parent / "test" / "_annotations.coco.json"
HPO_RECORD = Path(
    r"E:\PHD\Results\Quality Assessment\FINAL_B200"
    r"\session_20260618_113853\TwoClass\HPO_Config_009\hpo_record.json"
)
TRAIN_KWARGS = HPO_RECORD.parent / "run_meta" / "train_kwargs.json"
DEFAULT_MANUAL_LABELS = Path(
    r"C:\Users\SH37YE\OneDrive\Full_FOV_Master_Review_2026-06-02_1154.xlsx"
)
WD_ALIGN_LEFT = 0
WD_ALIGN_CENTER = 1
WD_BREAK_PAGE = 7
WD_AUTOFIT_CONTENT = 1
WD_FORMAT_DOCUMENT_DEFAULT = 16
WD_PAPER_A4 = 7
WD_STATISTIC_PAGES = 2


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--object-output", type=Path, default=OBJECT_OUTPUT)
    parser.add_argument("--calibration-output", type=Path, default=CALIBRATION_OUTPUT)
    parser.add_argument("--downstream-output", type=Path, default=None)
    parser.add_argument("--manual-labels", type=Path, default=DEFAULT_MANUAL_LABELS)
    parser.add_argument("--output-dir", type=Path, default=None)
    parser.add_argument("--split-summary", type=Path, default=SPLIT_SUMMARY)
    parser.add_argument("--test-coco", type=Path, default=TEST_COCO)
    parser.add_argument("--hpo-record", type=Path, default=HPO_RECORD)
    parser.add_argument("--train-kwargs", type=Path, default=None)
    parser.add_argument("--check-inputs", action="store_true", help="Validate artifact inputs without generating a report.")
    return parser.parse_args()


def read_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)
    if not isinstance(payload, dict):
        raise ValueError(f"Expected a JSON object in {path}")
    return payload


def validate_report_inputs(object_summary, confusion, calibration, downstream, rows):
    """Reject mixed operating points before writing any report artifacts."""
    selected = calibration["selected_thresholds"]
    settings = downstream["preflight"]["inference_settings"]
    for label, values in (
        ("held-out operating point", object_summary["operating_point"]["class_score_thresholds"]),
        ("held-out confusion matrix", confusion["class_score_thresholds"]),
        ("downstream inference", settings["class_score_thresholds"]),
    ):
        if set(values) != set(selected) or any(
            not math.isclose(float(values[key]), float(selected[key]), abs_tol=1e-6, rel_tol=0)
            for key in selected
        ):
            raise ValueError(f"{label} thresholds differ from independent calibration; regenerate the matching evaluation artifacts.")
    for value in (confusion["iou_threshold"], object_summary["confmat_iou"], object_summary["curve_iou"]):
        if not math.isclose(float(value), float(calibration["iou_threshold"]), abs_tol=1e-9):
            raise ValueError("Detection/calibration matching IoUs differ; do not combine these artifacts.")
    n = int(object_summary["images_processed"])
    count_rows = object_summary.get("count_error_metrics")
    if not count_rows or {r["class"] for r in count_rows} != set(selected):
        raise ValueError("Missing per-class count_error_metrics. Rerun eval_object_detection_RFDETR.py; downstream metrics cannot substitute for patch counts.")
    if len(count_rows) != len(selected) or any(int(r["n_images"]) != n for r in count_rows):
        raise ValueError("Count-error patch totals do not match the held-out evaluation.")
    if not rows or len(rows) != int(downstream["dataset"]["included_image_count"]):
        raise ValueError("Downstream CSV row count does not match its summary.")
    names = [r["source_image_name"].strip().casefold() for r in rows]
    if len(set(names)) != len(names):
        raise ValueError("Duplicate FOV names in downstream_predictions.csv.")
    # Paths are recorded by the producing machine; compare normalized stored paths.
    checkpoints = [object_summary.get("checkpoint"), calibration.get("checkpoint"),
                   downstream["preflight"].get("checkpoint")]
    checkpoint_ids = {str(p).replace("\\", "/").casefold() for p in checkpoints if p}
    if len(checkpoint_ids) > 1:
        raise ValueError("Artifacts record different checkpoints. Supply evaluations from the same selected checkpoint.")
    metadata = object_summary.get("count_error_analysis", {})
    if metadata and (metadata.get("unit") != "test_patch" or metadata.get("confidence_interval_method") is not None):
        raise ValueError("Unsupported count-error analysis metadata; this report expects descriptive patch-level point estimates.")
    if metadata and metadata.get("class_score_thresholds") != object_summary["operating_point"]["class_score_thresholds"]:
        raise ValueError("Count-error thresholds differ from the held-out operating point.")


def recover_patch_counts(summary: dict[str, Any], annotations: Path, predictions: Path):
    """Rebuild descriptive counts from the saved, unthresholded test detections.

    Older evaluations did not save count_error_metrics. Reuse their predictions,
    never downstream FOV counts, and require a complete, identifiable test set.
    """
    coco = read_json(annotations)
    detections = json.loads(predictions.read_text(encoding="utf-8"))
    image_ids = [image["id"] for image in coco["images"]]
    if (len(set(image_ids)) != len(image_ids)
            or len(image_ids) != int(summary["images_processed"])
            or summary.get("images_missing", 0)):
        raise ValueError("Cannot reconstruct patch counts for an incomplete test evaluation; rerun the object evaluator.")
    categories = {c["id"]: c["name"] for c in coco["categories"]}
    thresholds = summary["operating_point"]["class_score_thresholds"]
    if set(categories.values()) != set(thresholds):
        raise ValueError("Test categories differ from calibrated detector classes.")
    if float(summary["score_floor"]) > min(map(float, thresholds.values())):
        raise ValueError("Saved prediction score floor exceeds a calibrated threshold.")
    valid_ids = set(image_ids)
    annotated, predicted = Counter(), Counter()
    for item in coco["annotations"]:
        if int(item.get("iscrowd", 0)) != 0:
            continue
        if item["image_id"] not in valid_ids or item["category_id"] not in categories:
            raise ValueError("Unknown image or category in test annotations.")
        annotated[item["image_id"], item["category_id"]] += 1
    for item in detections:
        if item["image_id"] not in valid_ids or item["category_id"] not in categories:
            raise ValueError("Saved predictions do not belong to this test set.")
        score = float(item["score"])
        if not math.isfinite(score):
            raise ValueError("Non-finite saved prediction score.")
        # Match the evaluator's float32 threshold mask.
        if np.float32(score) >= np.float32(thresholds[categories[item["category_id"]]]):
            predicted[item["image_id"], item["category_id"]] += 1
    rows, metrics = [], []
    for category_id, name in categories.items():
        errors = []
        for image_id in image_ids:
            actual = annotated[image_id, category_id]
            estimate = predicted[image_id, category_id]
            error = estimate - actual
            errors.append(error)
            rows.append(dict(image_id=image_id, **{"class": name}, annotated_count=actual,
                             predicted_count=estimate, signed_error=error, absolute_error=abs(error)))
        metrics.append({"class": name, "n_images": len(errors),
                        "mean_absolute_error": float(np.mean(np.abs(errors))),
                        "median_absolute_error": float(np.median(np.abs(errors))),
                        "mean_signed_error": float(np.mean(errors))})
    previous = summary.get("count_error_metrics")
    metadata = summary.get("count_error_analysis", {})
    if metadata.get("confidence_interval_method") is not None:
        raise ValueError("This report only supports descriptive count errors; saved confidence intervals require an explicit reporting choice.")
    if previous:
        by_class = {row["class"]: row for row in previous}
        for row in metrics:
            if row["class"] not in by_class or any(
                not math.isclose(float(value), float(by_class[row["class"]][key]), abs_tol=1e-9)
                for key, value in row.items() if key != "class"
            ):
                raise ValueError("Saved count metrics disagree with the test predictions.")
    summary["count_error_metrics"] = metrics
    summary["count_error_analysis"] = {
        "unit": "test_patch", "class_score_thresholds": thresholds,
        "signed_error_definition": "predicted_minus_annotated",
        "confidence_interval_method": None, "bootstrap_replicates": None,
        "source": "reconstructed_from_saved_test_predictions",
    }
    return rows


def artifact_manifest(paths, calibration, object_summary):
    return {
        "schema_version": 2,
        "stages": ["coco_ranking", "independent_calibration", "calibrated_test_detection", "downstream_count_based_classification"],
        "reference_schemes": ["geckler", "collapsed_geckler", "murray_washington"],
        "calibrated_thresholds": calibration["selected_thresholds"],
        "test_patch_count": object_summary["images_processed"],
        "count_error_confidence_intervals": None,
        "inputs": {name: {"path": str(path.resolve()), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
                   for name, path in paths.items()},
        "report_builder_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
    }


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


from quality_reference_schemes import (
    normalized_count_bin as normalize_count_bin,
    count_bin_from_integer as integer_count_bin,
    geckler_class, collapsed_geckler_label, murray_washington_label as mw_label,
)
GECKLER_LABELS = [f"G{index}" for index in range(1, 7)]
MW_LABELS = ["Acceptable", "Unacceptable"]
COLLAPSED_GECKLER_LABELS = ["Acceptable", "Unacceptable", "Unknown"]


def count_adjustment_to_bin(predicted_count: Any, expert_bin: str) -> int:
    """Minimum signed change needed for a predicted count to enter an expert bin."""
    predicted = int(predicted_count)
    lower, upper = {
        "0-9": (0, 9),
        "10-25": (10, 25),
        "26+": (26, None),
    }[expert_bin]
    if predicted < lower:
        return lower - predicted
    if upper is not None and predicted > upper:
        return upper - predicted
    return 0


def signed_adjustment(value: int) -> str:
    return f"+{value}" if value > 0 else str(value)


def read_expert_references(path: Path, sheet_name: str = "Master") -> dict[str, dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        references: dict[str, dict[str, str]] = {}
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, min_col=1, max_col=6, values_only=True), start=2
        ):
            image, status, epithelial, leucocyte, annotator, comment = row
            if not image:
                continue
            if str(status or "").strip().casefold() == "er ikke i projektet":
                continue
            if epithelial is None and leucocyte is None:
                continue
            key = str(image).strip().casefold()
            if key in references:
                raise ValueError(f"Duplicate expert image name at workbook row {row_number}: {image}")
            epi_bin = normalize_count_bin(epithelial)
            leu_bin = normalize_count_bin(leucocyte)
            references[str(image).strip().casefold()] = {
                "epithelial_bin": epi_bin,
                "leucocyte_bin": leu_bin,
                "geckler_class": geckler_class(epi_bin, leu_bin),
                "mw_label": mw_label(epi_bin, leu_bin),
                "annotator": str(annotator or "").strip(),
                "comment": str(comment or "").strip(),
            }
        return references
    finally:
        workbook.close()


def named_classification_analysis(
    rows: Sequence[dict[str, str]],
    reference_key: str,
    predicted_key: str,
    labels: Sequence[str],
) -> dict[str, Any]:
    true = [row[reference_key] for row in rows]
    predicted = [row[predicted_key] for row in rows]
    matrix = np.zeros((len(labels), len(labels)), dtype=int)
    positions = {label: index for index, label in enumerate(labels)}
    for actual, estimate in zip(true, predicted):
        matrix[positions[actual], positions[estimate]] += 1
    per_class = []
    for index, label in enumerate(labels):
        tp = int(matrix[index, index])
        fn = int(matrix[index, :].sum() - tp)
        fp = int(matrix[:, index].sum() - tp)
        tn = len(rows) - tp - fn - fp
        precision = tp / (tp + fp) if tp + fp else 0.0
        recall = tp / (tp + fn) if tp + fn else 0.0
        per_class.append({
            "label": label, "support": tp + fn, "tp": tp, "fp": fp, "fn": fn,
            "precision": precision, "recall": recall,
            "f1": 2 * precision * recall / (precision + recall) if precision + recall else 0.0,
            "specificity": tn / (tn + fp) if tn + fp else 0.0,
        })
    return {
        "labels": list(labels),
        "reference_counts": dict(Counter(true)),
        "predicted_counts": dict(Counter(predicted)),
        "metrics": {
            "accuracy": float(accuracy_score(true, predicted)),
            "balanced_accuracy": float(balanced_accuracy_score(true, predicted)),
            "macro_f1": float(f1_score(true, predicted, labels=list(labels), average="macro", zero_division=0)),
            "cohen_kappa": float(cohen_kappa_score(true, predicted, labels=list(labels))),
        },
        "per_class": per_class,
        "matrix": matrix.tolist(),
    }


def save_confusion_matrix_figure(analysis: dict[str, Any], scheme: str, path: Path, *, detection: bool = False) -> None:
    """Render counts with a shared count scale and explicit reference/model axes."""
    from PIL import Image, ImageDraw, ImageFont

    labels = analysis["labels"]
    matrix = np.asarray(analysis["matrix"], dtype=int)
    if matrix.shape != (len(labels), len(labels)):
        raise ValueError("Confusion matrix dimensions do not match its labels.")
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    def font(size: int) -> Any:
        return ImageFont.truetype(str(font_path), size) if font_path.is_file() else ImageFont.load_default(size=size)

    image = Image.new("RGB", (1500, 1250), "white")
    draw = ImageDraw.Draw(image)
    draw.text((750, 45), scheme, font=font(42), fill="black", anchor="mt")
    draw.text((870, 130), "Predicted class" if detection else "Model classification (predicted)", font=font(32), fill="black", anchor="mt")
    draw.text((35, 175), "Annotated class" if detection else "Expert reference", font=font(28), fill="black")
    draw.text((35, 212), "(actual)", font=font(28), fill="black")
    left, top, span = 390, 280, 930
    cell = span // len(labels)
    maximum = max(1, int(matrix.max()))
    for index, label in enumerate(labels):
        draw.text((left + (index + .5) * cell, top - 25), label, font=font(27), fill="black", anchor="mb")
        draw.text((left - 24, top + (index + .5) * cell), label, font=font(29), fill="black", anchor="rm")
        for column, value in enumerate(matrix[index]):
            intensity = int(value) / maximum
            color = tuple(round(a + intensity * (b - a)) for a, b in zip((239, 246, 252), (31, 78, 121)))
            x, y = left + column * cell, top + index * cell
            draw.rectangle((x, y, x + cell, y + cell), fill=color, outline="white", width=3)
            draw.text((x + cell / 2, y + cell / 2), str(value), font=font(42), fill="white" if intensity > .55 else "black", anchor="mm")
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(240, 240))


def save_calibration_figure(sweep_path: Path, calibration: dict[str, Any], path: Path) -> None:
    """Plot both calibration F1 curves and label their selected maxima."""
    from PIL import Image, ImageDraw, ImageFont
    classes = ["Leucocyte", "Squamous Epithelial Cell"]
    rows = read_csv(sweep_path)
    if {r["class"] for r in rows} != set(classes):
        raise ValueError("Calibration sweep must contain the two detector classes.")
    grouped = {}
    selected_rows = []
    for name in classes:
        unique = {}
        for row in (r for r in rows if r["class"] == name):
            threshold = float(row["threshold"])
            if threshold in unique and row != unique[threshold]:
                raise ValueError("Conflicting duplicate thresholds in calibration sweep.")
            unique[threshold] = row
        points = [unique[key] for key in sorted(unique)]
        for point in points:
            if any(not math.isfinite(float(point[k])) or not 0 <= float(point[k]) <= 1
                   for k in ("threshold", "precision", "recall", "f1")):
                raise ValueError("Invalid calibration curve values.")
        chosen = [r for r in points if math.isclose(float(r["threshold"]), float(calibration["selected_thresholds"][name]), abs_tol=1e-6)]
        if len(chosen) != 1 or not math.isclose(float(chosen[0]["f1"]), max(float(r["f1"]) for r in points), abs_tol=1e-8):
            raise ValueError("Selected threshold is missing or is not an F1 maximum in the saved sweep.")
        selected_rows.append(chosen[0])
        grouped[name] = points
    if not math.isclose(sum(float(r["f1"]) for r in selected_rows)/2,
                        float(calibration["selected_joint_metrics"]["macro_f1"]), abs_tol=1e-8):
        raise ValueError("Selected curve points do not reproduce saved joint macro F1.")
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    def font(size):
        return ImageFont.truetype(str(font_path), size) if font_path.is_file() else ImageFont.load_default(size=size)
    canvas = Image.new("RGB", (2400, 1600), "white")
    draw = ImageDraw.Draw(canvas)
    colors = ["#0072B2", "#D55E00"]
    left, top, width, height = 190, 330, 2070, 1040
    xy = lambda x,y: (left+float(x)*width, top+(1-float(y))*height)
    draw.text((1200,55),"Confidence-threshold selection on the calibration set",anchor="mt",font=font(48),fill="black")
    for tick in range(11):
        value=tick/10
        x,y=xy(value,value)
        draw.line((left,y,left+width,y),fill="#dddddd",width=2)
        draw.text((left-25,y),f"{value:.1f}",anchor="rm",font=font(36),fill="black")
        draw.line((x,top+height,x,top+height+12),fill="black",width=2)
        draw.text((x,top+height+30),f"{value:.1f}",anchor="mt",font=font(36),fill="black")
    draw.text((left-75,top-60),"F1 score",font=font(38),fill="black")
    draw.line((left,top,left,top+height,left+width,top+height),fill="black",width=3)
    for index,name in enumerate(classes):
        color=colors[index]
        draw.line([xy(r["threshold"],r["f1"]) for r in grouped[name]],fill=color,width=8)
        threshold=float(calibration["selected_thresholds"][name])
        score=float(selected_rows[index]["f1"])
        px,py=xy(threshold,score)
        for y in range(round(py)+16,top+height,26):
            draw.line((px,y,px,min(y+13,top+height)),fill=color,width=3)
        draw.ellipse((px-14,py-14,px+14,py+14),fill=color,outline="white",width=3)
        label_x = 330 if index==0 else 1480
        title="Leucocytes" if index==0 else "Squamous epithelial cells"
        draw.text((label_x,150),title,font=font(44),fill=color)
        draw.text((label_x,211),f"Maximum F1 = {score:.4f} at threshold {threshold:.2f}",font=font(36),fill="black")
        endpoint_x=label_x+260
        draw.line((endpoint_x,275,px,py-22),fill=color,width=3)
    draw.text((left+width/2,1500),"Confidence threshold",anchor="mt",font=font(44),fill="black")
    path.parent.mkdir(parents=True,exist_ok=True)
    canvas.save(path,dpi=(400,400))


def latest_downstream_output(root: Path) -> Path:
    candidates = sorted(
        (
            path
            for path in root.iterdir()
            if path.is_dir()
            and (path / "downstream_evaluation_summary.json").is_file()
        ),
        key=lambda path: path.name,
    )
    if not candidates:
        raise FileNotFoundError(f"No completed downstream output was found in {root}")
    return candidates[-1]


def rgb(red: int, green: int, blue: int) -> int:
    """Convert RGB components to the integer representation used by Word."""
    return int(red) + int(green) * 256 + int(blue) * 65536


def fmt(value: Any, digits: int = 3) -> str:
    if value is None:
        return "—"
    if isinstance(value, (int, np.integer)):
        return f"{int(value):,}"
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    return f"{number:.{digits}f}"


def pct(value: Any, digits: int = 1) -> str:
    return f"{100.0 * float(value):.{digits}f}%"


def threshold_text(thresholds: dict[str, Any]) -> str:
    return "; ".join(
        f"{class_name}={float(value):.2f}"
        for class_name, value in thresholds.items()
    )


def sample_id(image_name: str) -> str:
    match = re.search(r"Sample\s*(\d+)", image_name, re.IGNORECASE)
    return f"Sample{match.group(1)}" if match else image_name


def per_class_standard_coco(
    test_coco_path: Path,
    predictions_path: Path,
) -> list[dict[str, Any]]:
    try:
        from pycocotools.coco import COCO
        from pycocotools.cocoeval import COCOeval
    except ImportError:
        saved_rows = read_csv(predictions_path.parent / "per_class_iou_sweep.csv")
        grouped = defaultdict(dict)
        for row in saved_rows:
            grouped[row["class"]][round(float(row["iou_threshold"]), 2)] = row
        grid = [round(0.50 + i * 0.05, 2) for i in range(10)]
        result = []
        for name, values in grouped.items():
            if any(iou not in values for iou in grid):
                raise ValueError("Saved per-class IoU sweep lacks the standard COCO 0.50–0.95 grid.")
            result.append({"class": name,
                           "AP@50:95": float(np.mean([float(values[iou]["AP"]) for iou in grid])),
                           "AP@50": float(values[0.5]["AP"]),
                           "AP@75": float(values[0.75]["AP"]),
                           "AR@100": float(np.mean([float(values[iou]["AR"]) for iou in grid]))})
        return result

    with contextlib.redirect_stdout(io.StringIO()):
        coco = COCO(str(test_coco_path))
        detections = coco.loadRes(str(predictions_path))
        rows: list[dict[str, Any]] = []
        for category in coco.loadCats(coco.getCatIds()):
            evaluator = COCOeval(coco, detections, "bbox")
            evaluator.params.catIds = [category["id"]]
            evaluator.evaluate()
            evaluator.accumulate()
            evaluator.summarize()
            rows.append(
                {
                    "class": category["name"],
                    "AP@50:95": float(evaluator.stats[0]),
                    "AP@50": float(evaluator.stats[1]),
                    "AP@75": float(evaluator.stats[2]),
                    "AR@100": float(evaluator.stats[8]),
                }
            )
    return rows


def downstream_descriptive_analysis(rows: Sequence[dict[str, str]]) -> dict[str, Any]:
    fields = {
        "kept_before_cross_class": "n_predictions_kept_before_duplicate_suppression",
        "cross_class_suppressed": "n_cross_class_duplicates_suppressed",
        "final": "n_predictions_kept", "leucocyte": "n_leucocyte",
        "epithelial": "n_squamous_epithelial_cell",
    }
    totals = {name: sum(int(row[field]) for row in rows) for name, field in fields.items()}
    totals["raw"] = sum(int(row.get("n_predictions_after_sahi_merge", row.get("n_predictions_raw", 0))) for row in rows)
    return {"aggregate_detections": totals}


class WordReport:
    def __init__(self) -> None:
        if win32com is None:
            raise RuntimeError("The Microsoft Word COM backend is unavailable.")
        self.word = win32com.client.DispatchEx("Word.Application")
        self.word.Visible = False
        self.word.DisplayAlerts = 0
        self.document = self.word.Documents.Add()
        self.selection = self.word.Selection
        self._configure_document()

    def _configure_document(self) -> None:
        for section in self.document.Sections:
            section.PageSetup.PaperSize = WD_PAPER_A4
            section.PageSetup.TopMargin = 2.0 * 28.3464567
            section.PageSetup.BottomMargin = 1.8 * 28.3464567
            section.PageSetup.LeftMargin = 2.0 * 28.3464567
            section.PageSetup.RightMargin = 2.0 * 28.3464567

        styles = self.document.Styles
        normal = styles.Item("Normal")
        normal.Font.Name = "Aptos"
        normal.Font.Size = 10
        normal.ParagraphFormat.SpaceAfter = 6
        normal.ParagraphFormat.LineSpacing = 13

        style_settings = {
            "Title": (24, (31, 78, 121)),
            "Subtitle": (13, (75, 85, 99)),
            "Heading 1": (16, (31, 78, 121)),
            "Heading 2": (13, (46, 116, 181)),
            "Heading 3": (11, (55, 65, 81)),
            "Caption": (9, (75, 85, 99)),
        }
        for name, (size, color) in style_settings.items():
            style = styles.Item(name)
            style.Font.Name = "Aptos"
            style.Font.Size = size
            style.Font.Color = rgb(*color)
            if name.startswith("Heading"):
                style.Font.Bold = True
                style.ParagraphFormat.KeepWithNext = True

        for section in self.document.Sections:
            footer = section.Footers.Item(1)
            footer.Range.Text = (
                "RF-DETR microscopy quality assessment — combined results  |  "
            )
            footer.Range.Font.Name = "Aptos"
            footer.Range.Font.Size = 8
            footer.Range.Font.Color = rgb(100, 116, 139)
            footer.Range.ParagraphFormat.Alignment = WD_ALIGN_CENTER
            footer.PageNumbers.Add()

        try:
            self.document.BuiltInDocumentProperties("Title").Value = (
                "RF-DETR Object Detection, Calibration, and Downstream Results"
            )
            self.document.BuiltInDocumentProperties("Subject").Value = (
                "End-to-end evaluation of microscopy specimen quality assessment"
            )
        except Exception:
            pass

    def paragraph(
        self,
        text: str = "",
        style: str = "Normal",
        alignment: int = WD_ALIGN_LEFT,
        bold: bool = False,
        italic: bool = False,
    ) -> None:
        self.selection.Style = self.document.Styles.Item(style)
        self.selection.ParagraphFormat.Alignment = alignment
        self.selection.Font.Bold = bold or style.startswith("Heading")
        self.selection.Font.Italic = italic
        self.selection.TypeText(str(text))
        self.selection.TypeParagraph()
        self.selection.Font.Bold = False
        self.selection.Font.Italic = False
        self.selection.ParagraphFormat.Alignment = WD_ALIGN_LEFT

    def bullets(self, items: Iterable[str]) -> None:
        for item in items:
            self.paragraph(str(item), style="List Bullet")

    def page_break(self) -> None:
        self.selection.InsertBreak(WD_BREAK_PAGE)

    def table(
        self,
        headers: Sequence[str],
        rows: Sequence[Sequence[Any]],
        font_size: float = 9,
    ) -> None:
        table = self.document.Tables.Add(
            self.selection.Range,
            len(rows) + 1,
            len(headers),
        )
        table.Style = "Table Grid"
        table.AllowAutoFit = True
        table.AutoFitBehavior(WD_AUTOFIT_CONTENT)
        table.Rows.Item(1).HeadingFormat = True
        table.Rows.Item(1).Range.ParagraphFormat.KeepWithNext = bool(rows)
        table.Rows.Item(1).Range.Font.Bold = True
        table.Rows.Item(1).Range.Font.Color = rgb(255, 255, 255)
        table.Rows.Item(1).Shading.BackgroundPatternColor = rgb(31, 78, 121)

        for column, value in enumerate(headers, start=1):
            table.Cell(1, column).Range.Text = str(value)
        for row_index, values in enumerate(rows, start=2):
            table.Rows.Item(row_index).AllowBreakAcrossPages = False
            if row_index % 2 == 1:
                table.Rows.Item(row_index).Shading.BackgroundPatternColor = rgb(
                    238,
                    244,
                    250,
                )
            for column, value in enumerate(values, start=1):
                table.Cell(row_index, column).Range.Text = str(value)

        table.Range.Font.Name = "Aptos"
        table.Range.Font.Size = font_size
        self.selection.SetRange(table.Range.End, table.Range.End)
        self.selection.TypeParagraph()

    def image(
        self,
        path: Path,
        caption: str,
        width_inches: float = 6.25,
    ) -> None:
        if not path.is_file():
            raise FileNotFoundError(f"Report figure is missing: {path}")
        self.selection.Style = self.document.Styles.Item("Normal")
        self.selection.ParagraphFormat.Alignment = WD_ALIGN_CENTER
        shape = self.selection.InlineShapes.AddPicture(
            str(path.resolve()),
            False,
            True,
        )
        maximum_width = 72.0 * width_inches
        if shape.Width > maximum_width:
            shape.LockAspectRatio = True
            shape.Width = maximum_width
        shape.Range.ParagraphFormat.KeepWithNext = True
        self.selection.SetRange(shape.Range.End, shape.Range.End)
        self.selection.TypeParagraph()
        self.paragraph(caption, style="Caption", alignment=WD_ALIGN_CENTER)

    def add_contents(self) -> None:
        self.paragraph("Contents", style="Heading 1")
        table_of_contents = self.document.TablesOfContents.Add(
            Range=self.selection.Range,
            UseHeadingStyles=True,
            UpperHeadingLevel=1,
            LowerHeadingLevel=3,
            UseHyperlinks=True,
        )
        self.selection.SetRange(
            table_of_contents.Range.End,
            table_of_contents.Range.End,
        )
        self.selection.TypeParagraph()
        self.page_break()

    def save(self, path: Path) -> dict[str, int]:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.document.Fields.Update()
        for toc in self.document.TablesOfContents:
            toc.Update()
        self.document.Repaginate()
        self.document.SaveAs2(
            str(path.resolve()),
            FileFormat=WD_FORMAT_DOCUMENT_DEFAULT,
        )
        statistics = {
            "pages": int(self.document.ComputeStatistics(WD_STATISTIC_PAGES)),
            "paragraphs": int(self.document.Paragraphs.Count),
            "tables": int(self.document.Tables.Count),
            "figures": int(self.document.InlineShapes.Count),
        }
        self.document.Close(SaveChanges=False)
        self.word.Quit()
        return statistics

    def close_without_saving(self) -> None:
        try:
            self.document.Close(SaveChanges=False)
        finally:
            self.word.Quit()


if win32com is None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor

    class WordReport:  # type: ignore[no-redef]
        """Portable python-docx backend used when local Word automation is absent."""

        def __init__(self) -> None:
            self.document = Document()
            section = self.document.sections[0]
            section.page_width = Cm(21.0)
            section.page_height = Cm(29.7)
            section.top_margin = Cm(2.0)
            section.bottom_margin = Cm(1.8)
            section.left_margin = Cm(2.0)
            section.right_margin = Cm(2.0)
            styles = self.document.styles
            styles["Normal"].font.name = "Aptos"
            styles["Normal"].font.size = Pt(10)
            for name, size, color in (
                ("Title", 24, (31, 78, 121)),
                ("Subtitle", 13, (75, 85, 99)),
                ("Heading 1", 16, (31, 78, 121)),
                ("Heading 2", 13, (46, 116, 181)),
                ("Heading 3", 11, (55, 65, 81)),
                ("Caption", 9, (75, 85, 99)),
            ):
                styles[name].font.name = "Aptos"
                styles[name].font.size = Pt(size)
                styles[name].font.color.rgb = RGBColor(*color)
            footer = section.footer.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer.add_run("RF-DETR microscopy quality assessment — combined results  |  ")
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            footer._p.append(field)
            self._figure_count = 0

        def paragraph(self, text: str = "", style: str = "Normal", alignment: int = 0,
                      bold: bool = False, italic: bool = False) -> None:
            paragraph = self.document.add_paragraph(style=style)
            paragraph.alignment = alignment
            run = paragraph.add_run(str(text))
            run.bold = bold or style.startswith("Heading")
            run.italic = italic

        def bullets(self, items: Iterable[str]) -> None:
            for item in items:
                self.paragraph(str(item), style="List Bullet")

        def page_break(self) -> None:
            self.document.add_page_break()

        def table(self, headers: Sequence[str], rows: Sequence[Sequence[Any]], font_size: float = 9) -> None:
            table = self.document.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            header_repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(header_repeat)
            for index, value in enumerate(headers):
                cell = table.rows[0].cells[index]
                cell.text = str(value)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.paragraphs[0].paragraph_format.keep_with_next = bool(rows)
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "1F4E79")
                cell._tc.get_or_add_tcPr().append(shading)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            for row_index, values in enumerate(rows):
                cells = table.add_row().cells
                for column, value in enumerate(values):
                    cells[column].text = str(value)
                    if row_index % 2:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "EEF4FA")
                        cells[column]._tc.get_or_add_tcPr().append(shading)
            for row in table.rows:
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = "Aptos"
                            run.font.size = Pt(font_size)
            if len(rows) <= 12:
                for row in table.rows[:-1]:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.keep_with_next = True
            self.document.add_paragraph()

        def image(self, path: Path, caption: str, width_inches: float = 6.25) -> None:
            if not path.is_file():
                raise FileNotFoundError(f"Report figure is missing: {path}")
            paragraph = self.document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(path.resolve()), width=Inches(width_inches))
            self.paragraph(caption, style="Caption", alignment=WD_ALIGN_PARAGRAPH.CENTER)
            self._figure_count += 1

        def add_contents(self) -> None:
            self.paragraph("Contents", style="Heading 1")
            self._contents_anchor = self.document.add_paragraph()
            self.page_break()

        def save(self, path: Path) -> dict[str, int]:
            path.parent.mkdir(parents=True, exist_ok=True)
            # A Word TOC field has no cached content in python-docx output.
            # Supply a readable outline without requiring manual field updates.
            for paragraph in list(self.document.paragraphs):
                if paragraph.style.name == "Heading 1" and paragraph.text != "Contents":
                    self._contents_anchor.insert_paragraph_before(paragraph.text)
            self.document.save(path)
            return {
                "pages": 0,
                "paragraphs": len(self.document.paragraphs),
                "tables": len(self.document.tables),
                "figures": self._figure_count,
            }

        def close_without_saving(self) -> None:
            return


def build_report(args: argparse.Namespace) -> tuple[Path, dict[str, int]]:
    object_output = args.object_output.resolve()
    calibration_output = args.calibration_output.resolve()
    downstream_output = (
        args.downstream_output.resolve()
        if args.downstream_output
        else latest_downstream_output(DOWNSTREAM_ROOT).resolve()
    )
    if args.output_dir:
        output_dir = args.output_dir.resolve()
    else:
        output_dir = (
            PROJECT_ROOT
            / "EvaluationOutput"
            / f"Combined_RFDETR_Study_Report_{datetime.now():%Y%m%d-%H%M%S}"
        )
    report_path = output_dir / "RFDETR_Combined_Study_Results.docx"

    input_paths = {
        "object_summary": object_output / "eval_summary.json",
        "object_confusion": object_output / "confusion_matrix.json",
        "object_predictions": object_output / "predictions_coco.json",
        "object_per_class_iou": object_output / "per_class_iou_sweep.csv",
        "calibration": calibration_output / "threshold_calibration_summary.json",
        "calibration_sweep": calibration_output / "per_class_threshold_sweep.csv",
        "downstream_summary": downstream_output / "downstream_evaluation_summary.json",
        "downstream_predictions": downstream_output / "downstream_predictions.csv",
        "expert_counts": args.manual_labels,
        "split_summary": args.split_summary,
        "test_annotations": args.test_coco,
        "hpo_record": args.hpo_record,
        "train_kwargs": args.train_kwargs or args.hpo_record.parent / "run_meta" / "train_kwargs.json",
    }
    missing = [f"{name}: {path}" for name, path in input_paths.items() if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing report inputs; supply the corresponding path options:\n" + "\n".join(missing))

    object_summary = read_json(object_output / "eval_summary.json")
    object_confusion = read_json(object_output / "confusion_matrix.json")
    calibration = read_json(calibration_output / "threshold_calibration_summary.json")
    downstream = read_json(downstream_output / "downstream_evaluation_summary.json")
    split = read_json(args.split_summary)
    hpo = read_json(args.hpo_record)
    train_kwargs_path = args.train_kwargs or args.hpo_record.parent / "run_meta" / "train_kwargs.json"
    train_kwargs = read_json(train_kwargs_path)
    downstream_rows = read_csv(downstream_output / "downstream_predictions.csv")
    patch_count_rows = recover_patch_counts(object_summary, args.test_coco, input_paths["object_predictions"])
    # Independent reconciliation against the saved confusion-matrix margins.
    matrix = np.asarray(object_confusion["matrix"])
    for index, name in enumerate(object_confusion["labels"][:-1]):
        counts = [r for r in patch_count_rows if r["class"] == name]
        if (sum(r["annotated_count"] for r in counts) != int(matrix[index].sum())
                or sum(r["predicted_count"] for r in counts) != int(matrix[:, index].sum())):
            raise ValueError("Recovered patch counts disagree with saved confusion-matrix totals.")
    validate_report_inputs(object_summary, object_confusion, calibration, downstream, downstream_rows)
    expert_references = read_expert_references(args.manual_labels.resolve())
    if set(expert_references) != {r["source_image_name"].strip().casefold() for r in downstream_rows}:
        raise ValueError("Eligible expert FOVs differ from saved predictions; rerun downstream inference for the current cohort.")
    for row in downstream_rows:
        key = row["source_image_name"].strip().casefold()
        if key not in expert_references:
            raise ValueError(
                f"No revised expert reference was found for {row['source_image_name']!r}"
            )
        reference = expert_references[key]
        row["expert_epithelial_bin"] = reference["epithelial_bin"]
        row["expert_leucocyte_bin"] = reference["leucocyte_bin"]
        row["expert_geckler_class"] = reference["geckler_class"]
        row["expert_collapsed_geckler_label"] = collapsed_geckler_label(
            reference["geckler_class"]
        )
        row["expert_mw_label"] = reference["mw_label"]
        predicted_epi_bin = integer_count_bin(row["n_squamous_epithelial_cell"])
        predicted_leu_bin = integer_count_bin(row["n_leucocyte"])
        row["predicted_geckler_class"] = geckler_class(
            predicted_epi_bin, predicted_leu_bin
        )
        row["predicted_collapsed_geckler_label"] = collapsed_geckler_label(
            row["predicted_geckler_class"]
        )
        row["predicted_mw_label"] = mw_label(predicted_epi_bin, predicted_leu_bin)

    geckler_analysis = named_classification_analysis(
        downstream_rows,
        "expert_geckler_class",
        "predicted_geckler_class",
        GECKLER_LABELS,
    )
    collapsed_geckler_analysis = named_classification_analysis(
        downstream_rows,
        "expert_collapsed_geckler_label",
        "predicted_collapsed_geckler_label",
        COLLAPSED_GECKLER_LABELS,
    )
    mw_analysis = named_classification_analysis(
        downstream_rows,
        "expert_mw_label",
        "predicted_mw_label",
        MW_LABELS,
    )
    scheme_analyses = {
        "Geckler (six groups)": geckler_analysis,
        "Collapsed Geckler (three-way)": collapsed_geckler_analysis,
        "Murray-Washington (binary)": mw_analysis,
    }
    if args.check_inputs:
        print("Report inputs validated; no report was written.")
        return report_path, {"checked_only": 1}
    output_dir.mkdir(parents=True, exist_ok=True)
    calibration_figure = output_dir / "supplementary_threshold_calibration.png"
    save_calibration_figure(input_paths["calibration_sweep"], calibration, calibration_figure)
    calibration_caption = (
        "Independent confidence-threshold calibration. F1 curves are shown "
        f"for leucocytes and squamous epithelial cells on {calibration['image_count']} calibration patches "
        f"from {calibration['task_count']} specimens, with within-class prediction-to-annotation matching at IoU {calibration['iou_threshold']:.2f}. "
        "Dots and labels identify the maximum F1 scores and selected thresholds "
        f"(leucocytes: F1 {calibration['selected_joint_metrics']['leucocyte_f1']:.4f} at {calibration['selected_thresholds']['Leucocyte']:.2f}; "
        f"squamous epithelial cells: F1 {calibration['selected_joint_metrics']['epithelial_f1']:.4f} at {calibration['selected_thresholds']['Squamous Epithelial Cell']:.2f}). "
        "Dashed lines project these thresholds onto the horizontal axis. "
        "The threshold pair maximized the unweighted mean of the two "
        "class-specific F1 scores (macro F1). Curves describe the calibration set used "
        "for threshold selection, not independent test-set performance."
    )
    (output_dir / "supplementary_threshold_calibration_caption.txt").write_text(calibration_caption, encoding="utf-8")
    (output_dir / "patch_count_error_results.json").write_text(
        json.dumps({"analysis": object_summary["count_error_analysis"],
                    "metrics": object_summary["count_error_metrics"]}, indent=2), encoding="utf-8")
    clean_fields = ["source_image_name", "n_leucocyte", "n_squamous_epithelial_cell",
                    "expert_epithelial_bin", "expert_leucocyte_bin", "expert_geckler_class",
                    "expert_collapsed_geckler_label", "expert_mw_label", "predicted_geckler_class",
                    "predicted_collapsed_geckler_label", "predicted_mw_label"]
    for filename, records, fields in (
        ("per_image_count_errors.csv", patch_count_rows, list(patch_count_rows[0])),
        ("count_based_fov_predictions.csv", downstream_rows, clean_fields),
    ):
        with (output_dir / filename).open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(records)
    manifest = artifact_manifest(input_paths, calibration, object_summary)
    manifest_path = output_dir / "report_provenance.json"
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    (output_dir / "count_based_classification_results.json").write_text(
        json.dumps(scheme_analyses, indent=2), encoding="utf-8"
    )
    # Generate from the validated matrix data; do not reuse a potentially stale PNG.
    object_matrix_path = output_dir / "calibrated_test_detection_confusion_matrix.png"
    save_confusion_matrix_figure(object_confusion, "Held-out detection at calibrated thresholds", object_matrix_path, detection=True)
    per_class_coco = per_class_standard_coco(
        args.test_coco,
        object_output / "predictions_coco.json",
    )
    n_clusters = len({sample_id(row["source_image_name"]) for row in downstream_rows})
    descriptive = downstream_descriptive_analysis(downstream_rows)

    report = WordReport()
    try:
        report.paragraph(
            "End-to-End Evaluation of RF-DETR-Based\n"
            "Microscopy Specimen Quality Assessment",
            style="Title",
            alignment=WD_ALIGN_CENTER,
        )
        report.paragraph(
            "Object detection • independent threshold calibration • "
            "final expert-referenced downstream evaluation",
            style="Subtitle",
            alignment=WD_ALIGN_CENTER,
        )
        report.paragraph(
            f"Generated {datetime.now():%d %B %Y}",
            alignment=WD_ALIGN_CENTER,
        )
        report.paragraph("")
        report.table(
            ["Item", "Locked value"],
            [
                ["Detector", "RFDETR2XLarge, HPO configuration 009, EMA checkpoint"],
                ["Internal model resolution", "880 × 880 pixels"],
                ["SAHI full-FOV slice size", "640 × 640 pixels with 20% overlap"],
                [
                    "Class confidence thresholds",
                    threshold_text(calibration["selected_thresholds"]),
                ],
                ["Final downstream set", f"{len(downstream_rows)} FOVs from {n_clusters} samples"],
                ["Artifact provenance", str(manifest_path.resolve())],
            ],
        )
        report.paragraph("Headline outcomes", style="Heading 2")
        report.table(
            ["Evaluation stage", "Primary result"],
            [
                [
                    "Held-out object detection",
                    "COCO AP@50:95 "
                    + fmt(object_summary["coco_standard"]["AP@50:95"], 4)
                    + "; AP@50 "
                    + fmt(object_summary["coco_standard"]["AP@50"], 4),
                ],
                [
                    "Independent calibration",
                    "Macro F1 "
                    + fmt(calibration["selected_joint_metrics"]["macro_f1"], 4)
                    + " at " + threshold_text(calibration["selected_thresholds"]),
                ],
                [
                    "Final downstream task",
                    "Geckler accuracy "
                    + pct(geckler_analysis["metrics"]["accuracy"], 2)
                    + "; collapsed Geckler accuracy "
                    + pct(collapsed_geckler_analysis["metrics"]["accuracy"], 2)
                    + "; Murray-Washington accuracy "
                    + pct(mw_analysis["metrics"]["accuracy"], 2),
                ],
            ],
        )
        report.page_break()
        report.add_contents()

        report.paragraph("1. Executive summary", style="Heading 1")
        report.paragraph(
            "This report consolidates the complete three-stage evaluation of the "
            "final two-class RF-DETR model. The detector was first assessed on a "
            "held-out object-detection test split. Class-specific confidence "
            "thresholds were then selected on a separate calibration set. Those "
            "thresholds and all downstream inference settings were locked before "
            "the expert-referenced downstream FOV labels were analyzed."
        )
        report.bullets(
            [
                "The detector achieved COCO AP@50:95 of "
                + fmt(object_summary["coco_standard"]["AP@50:95"], 4)
                + " and AP@50 of "
                + fmt(object_summary["coco_standard"]["AP@50"], 4)
                + f" on {object_summary['images_processed']} held-out patches.",
                f"Independent calibration on {calibration['image_count']} patches selected "
                + threshold_text(calibration["selected_thresholds"]) + "; joint macro F1 was "
                + fmt(calibration["selected_joint_metrics"]["macro_f1"], 4)
                + ".",
                f"The final downstream evaluation included {len(downstream_rows)} FOVs "
                f"after {downstream['dataset']['excluded_image_count']} exclusions. The references "
                "were derived independently under full Geckler, collapsed Geckler, "
                "and Murray-Washington rules.",
                "No threshold or quality-rule parameter was changed after the "
                "downstream labels were inspected.",
            ]
        )

        report.paragraph("2. Evaluation design and data separation", style="Heading 1")
        report.paragraph(
            "The study design deliberately separates model fitting, standard "
            "detector testing, threshold calibration, and final downstream "
            "testing. This prevents downstream test performance from influencing "
            "the confidence thresholds used to decide whether detections enter "
            "the count-based quality rule."
        )
        counts = split["counts"]
        report.table(
            ["Partition/stage", "Samples", "Images", "Leucocyte", "Epithelial", "Purpose"],
            [
                [
                    "Detector training",
                    counts["train"]["n_samples"],
                    counts["train"]["n_images"],
                    counts["train"]["n_target_boxes_by_class"]["Leucocyte"],
                    counts["train"]["n_target_boxes_by_class"][
                        "Squamous Epithelial Cell"
                    ],
                    "Model fitting",
                ],
                [
                    "Detector validation",
                    counts["valid"]["n_samples"],
                    counts["valid"]["n_images"],
                    counts["valid"]["n_target_boxes_by_class"]["Leucocyte"],
                    counts["valid"]["n_target_boxes_by_class"][
                        "Squamous Epithelial Cell"
                    ],
                    "HPO/model selection",
                ],
                [
                    "Detector test",
                    counts["test"]["n_samples"],
                    counts["test"]["n_images"],
                    counts["test"]["n_target_boxes_by_class"]["Leucocyte"],
                    counts["test"]["n_target_boxes_by_class"][
                        "Squamous Epithelial Cell"
                    ],
                    "Held-out object detection",
                ],
                [
                    "Calibration",
                    calibration["task_count"],
                    calibration["image_count"],
                    calibration["annotation_counts"]["Leucocyte"],
                    calibration["annotation_counts"]["Squamous Epithelial Cell"],
                    "Threshold selection only",
                ],
                [
                    "Downstream test",
                    n_clusters,
                    downstream["dataset"]["included_image_count"],
                    "—",
                    "—",
                    "Final FOV quality evaluation",
                ],
            ],
        )
        report.paragraph(
            "Detector partitions are separated by source specimen. Downstream exclusions "
            "are listed individually in Appendix A; the cohort sizes above are read from the supplied results."
        )

        report.paragraph("3. Selected model and training settings", style="Heading 1")
        report.paragraph("3.1 Model selection and training", style="Heading 2")
        report.table(
            ["Parameter", "Value"],
            [
                ["Selected configuration", "HPO configuration 009"],
                ["Architecture", hpo["MODEL_CLS"]],
                ["Checkpoint", "checkpoint_best_ema.pth"],
                ["Internal resolution", hpo["RESOLUTION"]],
                ["Input data mode", hpo["input_mode"] + " pixel source patches"],
                ["Queries", hpo["NUM_QUERIES"]],
                ["Batch size / gradient accumulation", f"{hpo['BATCH']} / {hpo['GRAD_ACCUM_STEPS']}"],
                ["Learning rate", f"{hpo['LR']:.1e}"],
                ["Weight decay / dropout", f"{hpo['WEIGHT_DECAY']} / {hpo['DROPOUT']}"],
                ["Multi-scale training", str(hpo["MULTI_SCALE"])],
                ["Maximum epochs / best epoch", f"{hpo['EPOCHS']} / {hpo['best_epoch']}"],
                ["Validation AP@50", fmt(hpo["val_AP50"], 4)],
                ["Validation mAP@50:95", fmt(hpo["val_mAP5095"], 4)],
                ["Random seed", hpo["SEED"]],
                ["Patch size / attention windows", f"{train_kwargs['patch_size']} / {train_kwargs['num_windows']}"],
            ],
        )
        report.paragraph(
            "The microscopy patches remained 640 × 640 at the data interface. "
            "RF-DETR internally resized them to the selected 880 × 880 model "
            "resolution. The same distinction was retained for downstream SAHI: "
            "the slice is 640 × 640, while the model internally operates at 880."
        )

        report.paragraph("4. Held-out COCO ranking evaluation", style="Heading 1")
        report.paragraph(
            "Standard COCO ranking metrics were calculated from detections "
            f"retained at the numerical score floor of {object_summary['score_floor']}. "
            "They do not use the calibrated thresholds. Threshold-dependent held-out "
            "results are reported in Section 6, after independent calibration in Section 5."
        )
        overall = object_summary["coco_standard"]
        report.table(
            ["Metric", "Overall result"],
            [
                ["COCO AP@50:95", fmt(overall["AP@50:95"], 4)],
                ["AP@50", fmt(overall["AP@50"], 4)],
                ["AP@75", fmt(overall["AP@75"], 4)],
                ["AR@1", fmt(overall["AR@1"], 4)],
                ["AR@10", fmt(overall["AR@10"], 4)],
                ["AR@100", fmt(overall["AR@100"], 4)],
            ],
        )
        report.paragraph("Per-class standard COCO metrics", style="Heading 2")
        report.table(
            ["Class", "AP@50:95", "AP@50", "AP@75", "AR@100"],
            [
                [
                    row["class"],
                    fmt(row["AP@50:95"], 4),
                    fmt(row["AP@50"], 4),
                    fmt(row["AP@75"], 4),
                    fmt(row["AR@100"], 4),
                ]
                for row in per_class_coco
            ],
        )
        report.image(
            object_output / "map_by_iou_threshold.png",
            "Figure 1. Detector AP across IoU thresholds. The standard headline "
            "COCO result remains AP@50:95.",
        )
        report.paragraph("5. Independent confidence-threshold calibration", style="Heading 1")
        report.paragraph(
            f"Calibration used {calibration['image_count']} patches from {calibration['task_count']} tasks. "
            f"Predictions were matched at IoU={calibration['iou_threshold']:.2f}. "
            f"The threshold grid ranged from {calibration['threshold_grid']['min']:.2f} "
            f"to {calibration['threshold_grid']['max']:.2f} in steps of {calibration['threshold_grid']['step']:.2f}; "
            f"the selection criterion was {calibration['selection_metric']}."
        )
        selected = calibration["selected_thresholds"]
        per_class = calibration["best_per_class_by_f1"]
        report.table(
            ["Class", "Threshold", "TP", "FP", "FN", "Precision", "Recall", "F1", "Jaccard"],
            [
                [
                    class_name,
                    fmt(selected[class_name], 2),
                    per_class[class_name]["tp"],
                    per_class[class_name]["fp"],
                    per_class[class_name]["fn"],
                    fmt(per_class[class_name]["precision"], 4),
                    fmt(per_class[class_name]["recall"], 4),
                    fmt(per_class[class_name]["f1"], 4),
                    fmt(per_class[class_name]["jaccard"], 4),
                ]
                for class_name in ("Leucocyte", "Squamous Epithelial Cell")
            ],
        )
        joint = calibration["selected_joint_metrics"]
        report.table(
            ["Joint metric", "Result"],
            [
                ["TP / FP / FN", f"{joint['tp']} / {joint['fp']} / {joint['fn']}"],
                ["Micro precision", fmt(joint["micro_precision"], 4)],
                ["Micro recall", fmt(joint["micro_recall"], 4)],
                ["Micro F1", fmt(joint["micro_f1"], 4)],
                ["Macro precision", fmt(joint["macro_precision"], 4)],
                ["Macro recall", fmt(joint["macro_recall"], 4)],
                ["Macro F1", fmt(joint["macro_f1"], 4)],
                ["Macro Jaccard", fmt(joint["macro_jaccard"], 4)],
            ],
        )
        report.image(
            calibration_figure,
            "Figure 2. " + calibration_caption,
        )
        report.image(
            calibration_output / "joint_macro_f1_heatmap.png",
            "Figure 3. Joint macro-F1 surface for the two class-specific "
            "confidence thresholds.",
        )
        report.paragraph(
            "The calibrated thresholds were subsequently treated as fixed "
            "method parameters. They were applied to the threshold-dependent "
            "detector results and carried unchanged into downstream SAHI "
            "inference."
        )

        report.paragraph("6. Held-out detection at calibrated thresholds", style="Heading 1")
        operating = object_summary["operating_point"]
        report.paragraph("6.1 Detection precision recall and confusion matrix", style="Heading 2")
        report.table(
            ["Thresholds", "TP", "FP", "FN", "Precision", "Recall", "F1", "FP/image"],
            [
                [
                    threshold_text(operating["class_score_thresholds"]),
                    operating["tp"],
                    operating["fp"],
                    operating["fn"],
                    fmt(operating["precision"], 4),
                    fmt(operating["recall"], 4),
                    fmt(operating["f1"], 4),
                    fmt(operating["fp_per_image"], 3),
                ]
            ],
        )
        report.image(
            object_matrix_path,
            "Figure 4. Held-out object-detection confusion matrix at the locked "
            f"thresholds ({threshold_text(operating['class_score_thresholds'])}) and IoU={object_confusion['iou_threshold']:.2f}. Rows are "
            "ground truth; columns are predictions. Background indicates misses "
            "or unmatched detections.",
        )
        report.paragraph(
            "The confusion matrix matches boxes without requiring class agreement, "
            "whereas precision/recall/F1 match within class. Their diagonal and true-positive "
            "totals need not be identical. Both use the held-out patches and the calibrated thresholds."
        )
        report.paragraph("6.2 Patch-level count agreement", style="Heading 2")
        count_rows = object_summary.get("count_error_metrics")
        if not count_rows:
            raise ValueError("Missing count_error_metrics in eval_summary.json. Rerun the current object evaluator; do not substitute downstream metrics.")
        report.table(
            ["Class", "Patches", "MAE", "Median AE", "Mean signed error"],
            [[r["class"], r["n_images"], fmt(r["mean_absolute_error"], 2),
              fmt(r["median_absolute_error"], 2), fmt(r["mean_signed_error"], 2)] for r in count_rows],
        )
        report.paragraph(
            "Count errors are in cells per test patch at the calibrated thresholds. Signed error "
            "is predicted minus annotated count. These are descriptive point estimates; "
            "no confidence intervals were calculated by this count-error analysis."
        )
        report.paragraph(
            "The saved test threshold sweep is descriptive sensitivity analysis "
            "only. Its test-set maximum was not used to select or revise the "
            "operating thresholds and is intentionally not reported as a "
            "selected operating point."
        )


        report.paragraph("7. Full-FOV count-based classification", style="Heading 1")
        report.paragraph("7.1 Full-FOV inference settings", style="Heading 2")
        settings = downstream["preflight"]["inference_settings"]
        cross_class = settings["cross_class_duplicate_suppression"]
        report.table(
            ["Setting", "Locked value"],
            [
                ["SAHI slice", f"{settings['slice_width']} × {settings['slice_height']}"],
                [
                    "Slice overlap",
                    f"{pct(settings['overlap_width_ratio'], 0)} horizontal; {pct(settings['overlap_height_ratio'], 0)} vertical",
                ],
                ["Standard whole-image prediction", str(settings["perform_standard_prediction"])],
                [
                    "Class confidence thresholds",
                    threshold_text(settings["class_score_thresholds"]),
                ],
                [
                    "Within-class postprocess",
                    f"{settings['postprocess']['type']}, "
                    f"{settings['postprocess']['match_metric']}="
                    f"{settings['postprocess']['match_threshold']:.2f}, "
                    + ("class-agnostic" if settings['postprocess']['class_agnostic'] else "class-aware"),
                ],
                [
                    "Cross-class duplicate suppression",
                    f"IOS≥{cross_class['ios_threshold']:.2f}, "
                    f"IoU≥{cross_class['iou_threshold']:.2f}, "
                    f"area ratio≥{cross_class['area_ratio_threshold']:.2f}",
                ],
            ],
        )

        report.paragraph("7.2 Count-based classification rules", style="Heading 2")
        report.paragraph(
            "For the revised primary analysis, both expert count bins and model "
            "counts were converted by the same published cell-count boundaries. "
            "Geckler groups G1-G3 contain >25 epithelial cells with respectively "
            "0-9, 10-25, or >25 leukocytes; G4 contains 10-25 epithelial cells and "
            ">25 leukocytes; G5 contains 0-9 epithelial cells and >25 leukocytes; "
            "G6 contains the remaining low-epithelial/low-leukocyte combinations. "
            "A three-way collapsed Geckler endpoint grouped G4-G5 as Acceptable, "
            "G1-G3 as Unacceptable, and G6 as Unknown. "
            "The Murray-Washington binary analysis defined an FOV as Acceptable "
            "only when it contained 0-9 epithelial cells and >25 leukocytes."
        )


        report.paragraph("7.3 Cohort flow and runtime", style="Heading 2")
        report.table(
            ["Item", "Count/result"],
            [
                ["Workbook rows reviewed", downstream["preflight"]["workbook_rows"]],
                ["Eligible FOVs analyzed", downstream["dataset"]["included_image_count"]],
                ["Represented samples", n_clusters],
                ["Excluded FOVs", downstream["dataset"]["excluded_image_count"]],
                ["Total evaluation time", f"{downstream['runtime']['total_evaluation_seconds']:.1f} s"],
                ["Mean detector time per FOV", f"{downstream['runtime']['mean_inference_seconds_per_image']:.2f} s"],
            ],
        )
        report.paragraph(
            f"Results are available for {len(downstream_rows)} eligible FOVs. "
            "Overlay availability is recorded in the downstream output manifest."
        )

        report.paragraph("7.4 Count-based agreement", style="Heading 2")
        report.paragraph(
            "Expert-recorded epithelial-cell and leucocyte count bins define the references. "
            "The same three count-based schemes are applied to predicted cell counts. "
            "Accuracy, mean class recall, unweighted macro F1, and unweighted Cohen's kappa "
            "describe agreement across FOVs. Unknown is retained in collapsed Geckler."
        )
        report.table(
            ["Scheme", "Reference distribution", "Model distribution"],
            [
                [
                    scheme,
                    "; ".join(f"{label}: {analysis['reference_counts'].get(label, 0)}" for label in analysis["labels"]),
                    "; ".join(f"{label}: {analysis['predicted_counts'].get(label, 0)}" for label in analysis["labels"]),
                ]
                for scheme, analysis in scheme_analyses.items()
            ],
            font_size=8,
        )
        report.table(
            ["Evaluation scheme", "Accuracy", "Balanced accuracy", "Macro F1", "Cohen's kappa"],
            [
                [
                    scheme,
                    fmt(analysis["metrics"]["accuracy"], 4),
                    fmt(analysis["metrics"]["balanced_accuracy"], 4),
                    fmt(analysis["metrics"]["macro_f1"], 4),
                    fmt(analysis["metrics"]["cohen_kappa"], 4),
                ]
                for scheme, analysis in scheme_analyses.items()
            ],
        )

        report.paragraph("7.5 Count-based classification confusion matrices", style="Heading 2")
        for scheme, analysis in (
            ("Geckler", geckler_analysis),
            ("Collapsed Geckler", collapsed_geckler_analysis),
            ("Murray-Washington", mw_analysis),
        ):
            figure_path = output_dir / (re.sub(r"[^a-z0-9]+", "_", scheme.lower()) + "_confusion_matrix.png")
            save_confusion_matrix_figure(analysis, scheme, figure_path)
            report.image(
                figure_path,
                f"{scheme} confusion matrix. Cells show FOV counts (N = {sum(map(sum, analysis['matrix']))}); "
                "rows are expert references and columns are model predictions. Darker cells indicate larger counts.",
                width_inches=5.8,
            )
            report.paragraph(f"{scheme} reference (rows) versus model classification (columns)")
            report.table(
                ["Reference"] + analysis["labels"],
                [
                    [label] + values
                    for label, values in zip(analysis["labels"], analysis["matrix"])
                ],
                font_size=8,
            )

        report.paragraph("7.6 Detection counts after post-processing", style="Heading 2")
        aggregate = descriptive["aggregate_detections"]
        report.paragraph(
            f"Across all FOVs, {aggregate['raw']:,} predictions after SAHI merging were "
            f"reduced to {aggregate['kept_before_cross_class']:,} candidates by "
            "the locked confidence thresholds. "
            f"Cross-class suppression removed {aggregate['cross_class_suppressed']:,} "
            f"overlapping predictions, leaving {aggregate['final']:,} detections "
            f"({aggregate['leucocyte']:,} leucocytes and "
            f"{aggregate['epithelial']:,} epithelial cells)."
        )

        report.paragraph("8. Integrated interpretation", style="Heading 1")
        report.bullets(
            [
                "The final RFDETR2XLarge detector showed strong held-out "
                "localization performance, especially for squamous epithelial "
                "cells. Leucocyte detection remained the weaker component.",
                "Independent calibration produced high object-level F1 and "
                "provided a defensible, prespecified operating point for any "
                "threshold-dependent result.",
                "Downstream agreement under each reference scheme measures the entire chain: "
                "sliced full-FOV inference, duplicate handling, class-specific "
                "thresholding, count aggregation, and the selected published "
                "cell-count classification.",
                "Geckler preserves six count-pattern groups, whereas "
                "collapsed Geckler treats G4-G5 as acceptable, G1-G3 as unacceptable, "
                "and G6 as unknown. Murray-Washington uses the narrower G5-like "
                "acceptable-cellularity criterion.",
                "The final downstream set must remain untouched for optimization. "
                "Any future revision of thresholds or the count-to-quality rule "
                "should be developed on a new training/development cohort and "
                "then tested on an independent external set.",
            ]
        )

        report.paragraph("9. Limitations", style="Heading 1")
        report.bullets(
            [
                "The downstream reference is the final expert-review workbook. "
                "The experts recorded count bins rather than exact counts, so the "
                "analysis can resolve published threshold categories but cannot "
                "recover within-bin cell counts.",
                "Excluded images and their recorded reasons are listed in Appendix A.",
                "The downstream analysis contains multiple FOVs from some source "
                "samples. The reported point estimates are descriptive and do not "
                "treat those FOVs as independent specimens.",
                "Confidence thresholds were calibrated on 640 × 640 object-level "
                "patches, whereas the downstream task uses overlapping SAHI "
                "slices extracted from full FOVs. The locked approach preserves "
                "test independence but does not remove this deployment-domain "
                "difference.",
                "The full Geckler, collapsed Geckler, and Murray-Washington rules "
                "were applied identically to expert bins and model counts and were "
                "not tuned on downstream outcomes.",
                "Point estimates do not quantify specimen-level sampling uncertainty "
                "or establish external generalizability.",
            ]
        )

        report.paragraph("10. Reproducibility and result artifacts", style="Heading 1")
        report.table(
            ["Artifact", "Location"],
            [
                ["Source artifacts and checksums", str(manifest_path.resolve())],
                ["Object-detection summary", str((object_output / "eval_summary.json").resolve())],
                ["Object predictions", str((object_output / "predictions_coco.json").resolve())],
                ["Calibration summary", str((calibration_output / "threshold_calibration_summary.json").resolve())],
                ["Calibration predictions", str((calibration_output / "calibration_predictions.json").resolve())],
                ["Revised expert count workbook", str(args.manual_labels.resolve())],
                ["Downstream summary", str((downstream_output / "downstream_evaluation_summary.json").resolve())],
                ["Per-FOV count-based results", str((output_dir / "count_based_fov_predictions.csv").resolve())],
                ["Downstream exclusions", str((downstream_output / "downstream_exclusions.csv").resolve())],
                ["Downstream overlays", str((downstream_output / "overlays").resolve())],
            ],
            font_size=8,
        )
        report.paragraph(
            "The report was generated directly from the saved machine-readable "
            "artifacts above. All headline values can therefore be traced to a "
            "JSON or CSV result file."
        )

        report.page_break()
        report.paragraph("Appendix A. Downstream exclusions", style="Heading 1")
        exclusions = downstream["dataset"]["exclusions"]
        report.table(
            ["Workbook row", "Image", "Reason", "Comment"],
            [
                [
                    row["source_row"],
                    row["source_image_name"],
                    row["exclusion_reason"],
                    row["comment"] or "—",
                ]
                for row in exclusions
            ],
            font_size=8,
        )

        report.paragraph("Appendix B. Misclassified downstream FOVs", style="Heading 1")
        report.paragraph(
            "The tables list every disagreement under each endpoint so individual "
            "FOVs can be located by workbook row and filename. Expert bins are the "
            "revised manual epithelial/leukocyte categories; model counts are the "
            "detector-derived cell totals. The Epi / Leu difference is the minimum "
            "signed adjustment needed to move each model count into its expert bin: "
            "negative means fewer predicted cells are needed, positive means more, "
            "and zero means the count is already in the correct bin. Complete results "
            f"for all {len(downstream_rows)} FOVs remain available in count_based_fov_predictions.csv."
        )
        error_schemes = [
            (
                "B.1 Geckler classification",
                "expert_geckler_class",
                "predicted_geckler_class",
            ),
            (
                "B.2 Collapsed Geckler classification",
                "expert_collapsed_geckler_label",
                "predicted_collapsed_geckler_label",
            ),
            (
                "B.3 Murray-Washington classification",
                "expert_mw_label",
                "predicted_mw_label",
            ),
        ]
        for heading, reference_key, predicted_key in error_schemes:
            errors = [
                row
                for row in downstream_rows
                if row[reference_key] != row[predicted_key]
            ]
            heading = f"{heading} ({len(errors)} disagreements)"
            if heading.startswith(("B.2 ", "B.3 ")):
                report.page_break()
            report.paragraph(heading, style="Heading 2")
            report.table(
                [
                    "Workbook row",
                    "FOV filename",
                    "Expert bins (Epi; Leu)",
                    "Reference",
                    "Model counts (Epi; Leu)",
                    "Epi / Leu",
                    "Prediction",
                ],
                [
                    [
                        row["source_row"],
                        row["source_image_name"],
                        f"{row['expert_epithelial_bin']}; {row['expert_leucocyte_bin']}",
                        row[reference_key],
                        f"{row['n_squamous_epithelial_cell']}; {row['n_leucocyte']}",
                        " / ".join(
                            [
                                signed_adjustment(
                                    count_adjustment_to_bin(
                                        row["n_squamous_epithelial_cell"],
                                        row["expert_epithelial_bin"],
                                    )
                                ),
                                signed_adjustment(
                                    count_adjustment_to_bin(
                                        row["n_leucocyte"],
                                        row["expert_leucocyte_bin"],
                                    )
                                ),
                            ]
                        ),
                        row[predicted_key],
                    ]
                    for row in errors
                ],
                font_size=6.5,
            )

        statistics = report.save(report_path)
    except Exception:
        report.close_without_saving()
        raise

    return report_path, statistics


def validate_saved_report(path: Path, expected_statistics: dict[str, int]) -> None:
    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError(f"Report was not created correctly: {path}")

    required_phrases = [
        "Held-out COCO ranking evaluation",
        "Independent confidence-threshold calibration",
        "Full-FOV count-based classification",
        "Geckler",
        "Murray-Washington",
        "Held-out detection at calibrated thresholds",
    ]
    if win32com is None:
        document = Document(path)
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        missing = [phrase for phrase in required_phrases if phrase not in text]
        if missing:
            raise RuntimeError(f"Report validation failed; missing text: {missing}")
        if len(document.tables) != expected_statistics["tables"]:
            raise RuntimeError("Table count changed after reopening the report.")
        if len(document.inline_shapes) != expected_statistics["figures"]:
            raise RuntimeError("Figure count changed after reopening the report.")
        if any(term in text.casefold() for term in ("original expert", "partially qualified", "not qualified", "manual_label")):
            raise RuntimeError("Superseded expert quality comparison found in the report.")
        return

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    document = None
    try:
        document = word.Documents.Open(str(path.resolve()), ReadOnly=True)
        document.Repaginate()
        pages = int(document.ComputeStatistics(WD_STATISTIC_PAGES))
        text = document.Content.Text
        if any(term in text.casefold() for term in ("original expert", "partially qualified", "not qualified", "manual_label")):
            raise RuntimeError("Superseded expert quality comparison found in the report.")
        missing = [phrase for phrase in required_phrases if phrase not in text]
        if missing:
            raise RuntimeError(f"Report validation failed; missing text: {missing}")
        if int(document.Tables.Count) != expected_statistics["tables"]:
            raise RuntimeError("Table count changed after reopening the report.")
        if int(document.InlineShapes.Count) != expected_statistics["figures"]:
            raise RuntimeError("Figure count changed after reopening the report.")
    finally:
        if document is not None:
            document.Close(SaveChanges=False)
        word.Quit()


def main() -> int:
    args = parse_args()
    report_path, statistics = build_report(args)
    if statistics.get("checked_only"):
        return 0
    validate_saved_report(report_path, statistics)
    print(
        json.dumps(
            {
                "report": str(report_path.resolve()),
                "size_bytes": report_path.stat().st_size,
                **statistics,
                "validation": "passed",
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
